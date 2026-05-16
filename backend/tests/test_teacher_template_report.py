from __future__ import annotations

from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.app.main import app
from backend.app.core.database import SessionLocal
from backend.app.models.course import Course
from backend.app.models.obe_outcome import OBEOutcome
from backend.app.models.student import Student

TEMPLATE_PATH = "backend/templates/teacher_input_template.xlsx"
ROOT_DIR = Path(__file__).resolve().parents[2]


def _build_teacher_workbook() -> bytes:
    wb = Workbook()
    ws_raw = wb.active
    ws_raw.title = "示例班级A期末"
    ws_main = wb.create_sheet("示例班级A")
    ws_summary = wb.create_sheet("示例班级A课程目标平均分和达成度 ")
    ws_chart = wb.create_sheet("示例班级A课程考核结果分析图")

    ws_raw["A1"] = "2024-2025学年第二学期末成绩单"
    ws_raw["A2"] = "课程名称：示例课程A"
    ws_raw["F2"] = "代课教师：示例教师"
    ws_raw["A3"] = "专业班级：示例班级A    学生人数：2"
    headers = ["学号", "姓名", "选择题", "判断题", "填空题", "计算题", "分析题", "分析题", "分析题", "分析题", "综合题", "成绩"]
    sub_headers = [None, None, None, None, None, None, "1", "2", "3", "总分", None, None]
    for index, value in enumerate(headers, start=1):
        ws_raw.cell(4, index).value = value
    for index, value in enumerate(sub_headers, start=1):
        ws_raw.cell(5, index).value = value
    ws_raw.append(["9000001", "", 18, 10, 4, 8, 8, 9, 9, 26, 14, 80])
    ws_raw.append(["9000002", "", 14, 6, 5, 8, 9, 8, 5, 22, 6, 61])

    ws_main["A2"] = "2024-2025学年第二学期"
    ws_main["A3"] = "承担单位：计算机科学与技术系             课程名称：示例课程A"
    ws_main["A4"] = "任课教师：示例教师                         班级：示例班级A"
    main_headers = {
        3: "单选题\n20分",
        4: "判断题\n10分",
        5: "填空题10分",
        6: "计算题10分",
        7: "分析题30分",
        10: "得分",
        11: "综合题\n20分",
    }
    for col, value in main_headers.items():
        ws_main.cell(5, col).value = value
    ws_main["G6"] = 1
    ws_main["H6"] = 2
    ws_main["I6"] = 3
    ws_main["C8"] = 20
    ws_main["D8"] = 10
    ws_main["E8"] = 10
    ws_main["F8"] = 10
    ws_main["G8"] = 10
    ws_main["H8"] = 10
    ws_main["I8"] = 10
    ws_main["J8"] = 30
    ws_main["K8"] = 20
    ws_main["M8"] = "=C8"
    ws_main["O8"] = "=E8"
    ws_main["Q8"] = "=D8+F8+J8"
    ws_main["S8"] = "=K8"
    student_rows = [
        ["9000001", "", 18, 10, 4, 8, 8, 9, 9, 26, 14, 80, 18, 0.9, 4, 0.4, 44, 0.88, 14, 0.7, 0.72],
        ["9000002", "", 14, 6, 5, 8, 9, 8, 5, 22, 6, 61, 14, 0.7, 5, 0.5, 36, 0.72, 6, 0.3, 0.55],
    ]
    for row_index, values in enumerate(student_rows, start=9):
        for col_index, value in enumerate(values, start=1):
            ws_main.cell(row_index, col_index).value = value

    ws_summary["C1"] = "总分"
    ws_summary["D1"] = "平均分"
    ws_summary["E1"] = "达成度"
    summary_rows = {
        2: ("课程目标1", 20, 16.0, 0.80),
        6: ("课程目标2", 10, 4.5, 0.45),
        9: ("课程目标3", 50, 40.0, 0.80),
        14: ("课程目标4", 20, 10.0, 0.50),
    }
    for row, (label, full_score, avg_score, achievement) in summary_rows.items():
        ws_summary.cell(row, 1).value = label
        ws_summary.cell(row + (3 if row in {2, 14} else 2), 3).value = full_score
        ws_summary.cell(row + (3 if row in {2, 14} else 2), 4).value = avg_score
        ws_summary.cell(row + (3 if row in {2, 14} else 2), 5).value = achievement

    ws_chart.append(["课程目标", "预测值", "真实值"])
    ws_chart.append(["课程目标1", 0.65, 0.8])
    ws_chart.append(["课程目标2", 0.65, 0.45])
    ws_chart.append(["课程目标3", 0.65, 0.8])
    ws_chart.append(["课程目标4", 0.65, 0.5])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


def _build_simple_input_workbook() -> bytes:
    wb = Workbook()
    ws_info = wb.active
    ws_info.title = "基本信息"
    ws_structure = wb.create_sheet("试卷结构")
    ws_outcomes = wb.create_sheet("课程目标")
    ws_scores = wb.create_sheet("学生成绩")

    ws_info.append(["字段", "内容"])
    ws_info.append(["学年学期", "2024-2025学年第二学期"])
    ws_info.append(["院系", "计算机科学与技术系"])
    ws_info.append(["课程名称", "示例课程A"])
    ws_info.append(["任课教师", "示例教师"])
    ws_info.append(["班级", "示例班级A"])
    ws_info.append(["考试日期", "2025年6月22日"])
    ws_info.append(["课程目标达成阈值", 0.65])

    ws_structure.append(["题号", "题型", "小题号", "满分", "课程目标编号"])
    structure_rows = [
        [1, "选择题", "", 20, "CO1"],
        [2, "判断题", "", 10, "CO3"],
        [3, "填空题", "", 10, "CO2"],
        [4, "计算题", "", 10, "CO3"],
        [5, "分析题", "1", 10, "CO3"],
        [6, "分析题", "2", 10, "CO3"],
        [7, "分析题", "3", 10, "CO3"],
        [8, "综合题", "", 20, "CO4"],
    ]
    for row in structure_rows:
        ws_structure.append(row)

    ws_outcomes.append(["课程目标编号", "指标点", "课程目标说明", "达成阈值"])
    ws_outcomes.append(["CO1", "指标点1-1", "掌握示例课程基础知识。", 0.65])
    ws_outcomes.append(["CO2", "指标点2-1", "能够完成示例问题分析。", 0.65])
    ws_outcomes.append(["CO3", "指标点4-2", "能够完成示例方案设计。", 0.65])
    ws_outcomes.append(["CO4", "指标点5-2", "能够综合应用示例知识解决问题。", 0.65])

    ws_scores.append(["学号", "姓名", "选择题", "判断题", "填空题", "计算题", "分析题1", "分析题2", "分析题3", "综合题"])
    ws_scores.append(["9000001", "学生A", 18, 10, 4, 8, 8, 9, 9, 14])
    ws_scores.append(["9000002", "学生B", 14, 6, 5, 8, 9, 8, 5, 6])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


def _build_cross_target_simple_workbook() -> bytes:
    wb = Workbook()
    ws_info = wb.active
    ws_info.title = "基本信息"
    ws_structure = wb.create_sheet("试卷结构")
    ws_outcomes = wb.create_sheet("课程目标")
    ws_scores = wb.create_sheet("学生成绩")

    ws_info.append(["字段", "内容"])
    ws_info.append(["学年学期", "演示学期"])
    ws_info.append(["院系", "示例系"])
    ws_info.append(["课程名称", "示例课程B"])
    ws_info.append(["任课教师", "示例教师B"])
    ws_info.append(["班级", "示例班级B"])
    ws_info.append(["考试日期", "2026年1月12日"])
    ws_info.append(["课程目标达成阈值", 0.7])

    ws_structure.append(["题号", "题型", "小题号", "满分", "课程目标编号"])
    ws_structure.append([1, "选择题", "1", 10, "CO1"])
    ws_structure.append([2, "选择题", "2", 10, "CO2"])
    ws_structure.append([3, "综合题", "", 20, "CO3"])

    ws_outcomes.append(["课程目标编号", "指标点", "课程目标说明", "达成阈值"])
    ws_outcomes.append(["CO1", "指标点1-1", "掌握示例课程B的基本概念和质量意识。", ""])
    ws_outcomes.append(["CO2", "指标点2-1", "能够开展示例问题分析并建立清晰模型。", ""])
    ws_outcomes.append(["CO3", "指标点3-2", "能够综合完成示例方案分析与改进。", 0.8])

    ws_scores.append(["学号", "姓名", "选择题1", "选择题2", "综合题"])
    ws_scores.append(["9000101", "学生甲", 8, 6, 16])
    ws_scores.append(["9000102", "学生乙", 6, 3, 10])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


def _build_achievement_analysis_workbook() -> bytes:
    wb = Workbook()
    ws_info = wb.active
    ws_info.title = "课程基本信息"
    ws_calc = wb.create_sheet("计算表")
    ws_summary = wb.create_sheet("汇总表")
    ws_analysis = wb.create_sheet("分析表")
    wb.create_sheet("课程目标1统计")
    wb.create_sheet("课程目标2统计")
    wb.create_sheet("总课程目标统计")

    ws_info["C6"] = "课程代码"
    ws_info["D6"] = "DEMO-SYLLABUS"
    ws_info["C7"] = "课程名称"
    ws_info["D7"] = "示例课程C"
    ws_info["C8"] = "授课对象"
    ws_info["D8"] = "示例班级C"
    ws_info["C10"] = "学生人数"
    ws_info["D10"] = 2
    ws_info["C11"] = "期望达成标准"
    ws_info["D11"] = 65
    ws_info["G6"] = "课程目标1"
    ws_info["G7"] = "单选题"
    ws_info["H7"] = "判断题"
    ws_info["G8"] = "课程目标2"
    ws_info["G9"] = "算法分析题"

    ws_calc["D4"] = "序号"
    ws_calc["E4"] = "学号"
    ws_calc["F4"] = "姓名"
    ws_calc["G4"] = "班级"
    ws_calc["H3"] = "课程目标1"
    ws_calc["J3"] = "课程目标1"
    ws_calc["K3"] = "课程目标2"
    ws_calc["H4"] = "单选题"
    ws_calc["I4"] = "判断题"
    ws_calc["J4"] = "达成度(%)"
    ws_calc["K4"] = "算法分析题"
    ws_calc["L4"] = "达成度(%)"
    ws_calc["H5"] = 20
    ws_calc["I5"] = 10
    ws_calc["K5"] = 10
    ws_calc["H6"] = 20
    ws_calc["I6"] = 10
    ws_calc["K6"] = 10
    rows = [
        [1, "9000201", "学生甲", "示例班级C", 18, 9, 90, 10, 100],
        [2, "9000202", "学生乙", "示例班级C", 16, 8, 80, 5, 50],
    ]
    for row_index, row in enumerate(rows, start=7):
        for col_offset, value in enumerate(row, start=4):
            ws_calc.cell(row_index, col_offset).value = value

    ws_summary.append([None, None, None, "序号", "学号", "姓名", "班级", "课程目标1", "课程目标2", "综合总目标"])
    ws_summary.append([None, None, None, 1, "9000201", "学生甲", "示例班级C", 90, 100, 91])
    ws_summary.append([None, None, None, 2, "9000202", "学生乙", "示例班级C", 80, 50, 72])

    ws_analysis["E20"] = "课程目标"
    ws_analysis["F20"] = "评价方式及比例"
    ws_analysis["H20"] = "目标分值"
    ws_analysis["I20"] = "实际平均分(百分制)%"
    ws_analysis["J20"] = "分目标权重%"
    ws_analysis["K20"] = "目标达成度%"
    ws_analysis["E21"] = "课程目标1"
    ws_analysis["F21"] = "单选题"
    ws_analysis["G21"] = "20.0%"
    ws_analysis["H21"] = 20
    ws_analysis["I21"] = 85
    ws_analysis["J21"] = 30
    ws_analysis["K21"] = 85
    ws_analysis["F22"] = "判断题"
    ws_analysis["G22"] = "10.0%"
    ws_analysis["H22"] = 10
    ws_analysis["I22"] = 85
    ws_analysis["E23"] = "课程目标2"
    ws_analysis["F23"] = "算法分析题"
    ws_analysis["G23"] = "10.0%"
    ws_analysis["H23"] = 10
    ws_analysis["I23"] = 75
    ws_analysis["J23"] = 10
    ws_analysis["K23"] = 75
    ws_analysis["K27"] = 82.5

    for sheet_name, values in {
        "课程目标1统计": [("9000201", 90), ("9000202", 80)],
        "课程目标2统计": [("9000201", 100), ("9000202", 50)],
        "总课程目标统计": [("9000201", 91), ("9000202", 72)],
    }.items():
        ws = wb[sheet_name]
        ws.append(["学号", "达成度", "平均值", "期望值"])
        for student_no, achievement in values:
            ws.append([student_no, achievement, 80, 65])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


def test_simple_input_workbook_preview_computes_analysis_without_result_sheets() -> None:
    workbook_bytes = _build_simple_input_workbook()
    with TestClient(app) as client:
        response = client.post(
            "/api/v1/import/teacher-workbook-preview",
            files={
                "file": (
                    "简洁输入模板.xlsx",
                    workbook_bytes,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert response.status_code == 200
    payload = response.json()["data"]
    assert payload["recognized_sheets"] == ["基本信息", "试卷结构", "课程目标", "学生成绩"]
    assert payload["meta"]["course_name"] == "示例课程A"
    assert payload["meta"]["student_count_actual"] == 2
    assert payload["score_stats"]["average_score"] == 70.5
    assert payload["course_outcome_count"] == 4
    assert payload["question_group_count"] == 6
    assert any(item["name"] == "课程目标与试题支撑关系" for item in payload["input_requirements"])


def test_achievement_analysis_workbook_preview_supports_teacher_export_format() -> None:
    workbook_bytes = _build_achievement_analysis_workbook()
    with TestClient(app) as client:
        response = client.post(
            "/api/v1/import/teacher-workbook-preview",
            files={
                "file": (
                    "sjfx_1751105887802.xls",
                    workbook_bytes,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert response.status_code == 200
    payload = response.json()["data"]
    assert payload["meta"]["course_name"] == "示例课程C"
    assert payload["meta"]["class_name"] == "示例班级C"
    assert payload["score_stats"]["total_students"] == 2
    assert payload["score_stats"]["average_score"] == 83
    assert payload["course_outcome_count"] == 2
    assert payload["question_group_count"] == 3


def test_teacher_workbook_task_report_uses_syllabus_course_metadata_and_outcomes() -> None:
    workbook_bytes = _build_achievement_analysis_workbook()
    with TestClient(app) as client:
        db = SessionLocal()
        try:
            course = db.query(Course).filter(Course.course_code == "DEMO-SYLLABUS").first()
            if not course:
                course = Course(
                    course_code="DEMO-SYLLABUS",
                    course_name="示例课程C",
                    term="演示学期",
                    department="计算机科学与技术系",
                    major="示例专业",
                    credit=3,
                    owner="示例负责人",
                    description="教学大纲导入课程",
                )
                db.add(course)
                db.flush()
            db.query(OBEOutcome).filter(OBEOutcome.course_id == course.id).delete()
            descriptions = [
                "理解示例课程C的基础概念与结构表达。",
                "能够分析示例问题并选择合适的方法。",
                "能够设计示例方案解决应用问题。",
                "能够综合运用示例方法完成结果验证。",
            ]
            for index, description in enumerate(descriptions, start=1):
                db.add(
                    OBEOutcome(
                        course_id=course.id,
                        co_code=f"CO{index}",
                        co_name=f"课程目标{index}",
                        indicator=f"指标点{index}",
                        description=description,
                        threshold=0.65,
                    )
                )
            db.commit()
        finally:
            db.close()

        import_response = client.post(
            "/api/v1/import/teacher-workbook-task",
            files={
                "file": (
                    "示例达成度计算.xlsx",
                    workbook_bytes,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"exam_date": "2026-01-10"},
        )
        assert import_response.status_code == 200
        run_id = import_response.json()["data"]["run_overview"]["run"]["id"]

        export_response = client.get(f"/api/v1/analysis-runs/{run_id}/export/docx")

    assert export_response.status_code == 200
    document = Document(BytesIO(export_response.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "吕梁学院演示学期" in text
    assert "计算机科学与技术系" in table_text
    assert "理解示例课程C" in table_text
    assert "能够分析示例问题" in table_text
    assert "依据《示例课程C》教学大纲" in table_text
    assert "依据待补充" not in table_text


def test_simple_input_workbook_task_builds_run_from_raw_inputs_only() -> None:
    workbook_bytes = _build_simple_input_workbook()
    with TestClient(app) as client:
        response = client.post(
            "/api/v1/import/teacher-workbook-task",
            files={
                "file": (
                    "简洁输入模板.xlsx",
                    workbook_bytes,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

        assert response.status_code == 200
        run_id = response.json()["data"]["run_overview"]["run"]["id"]
        dashboard_response = client.get(f"/api/v1/analysis-runs/{run_id}/dashboard")

    assert dashboard_response.status_code == 200
    dashboard = dashboard_response.json()["data"]
    outcomes = {item["co_code"]: item for item in dashboard["course_outcomes"]}
    assert outcomes["CO1"]["full_score"] == 20
    assert outcomes["CO1"]["avg_score"] == 16
    assert outcomes["CO2"]["supporting_groups"] == ["填空题"]
    assert outcomes["CO3"]["supporting_groups"] == ["判断题", "计算题", "分析题"]
    assert outcomes["CO3"]["full_score"] == 50
    analysis_group = next(item for item in dashboard["question_groups"] if item["qgroup_name"] == "分析题")
    assert analysis_group["avg_score"] == 24
    assert analysis_group["achievement"] == 0.8


def test_simple_input_task_preserves_per_question_outcome_and_goal_text() -> None:
    workbook_bytes = _build_cross_target_simple_workbook()
    with TestClient(app) as client:
        response = client.post(
            "/api/v1/import/teacher-workbook-task",
            files={
                "file": (
                    "示例课程B简洁输入.xlsx",
                    workbook_bytes,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

        assert response.status_code == 200
        run_id = response.json()["data"]["run_overview"]["run"]["id"]
        dashboard_response = client.get(f"/api/v1/analysis-runs/{run_id}/dashboard")

    assert dashboard_response.status_code == 200
    dashboard = dashboard_response.json()["data"]
    outcomes = {item["co_code"]: item for item in dashboard["course_outcomes"]}

    assert set(outcomes) >= {"CO1", "CO2", "CO3"}
    assert outcomes["CO1"]["full_score"] == 10
    assert outcomes["CO2"]["full_score"] == 10
    assert outcomes["CO1"]["avg_score"] == 7
    assert outcomes["CO2"]["avg_score"] == 4.5
    assert outcomes["CO1"]["supporting_groups"] == ["选择题"]
    assert outcomes["CO2"]["supporting_groups"] == ["选择题"]
    assert outcomes["CO1"]["indicator"] == "指标点1-1"
    assert "示例问题分析" in outcomes["CO2"]["description"]
    assert outcomes["CO1"]["threshold"] == 0.7
    assert outcomes["CO3"]["threshold"] == 0.8


def test_simple_input_template_can_omit_goal_descriptions_after_syllabus_import() -> None:
    workbook = load_workbook(BytesIO(_build_cross_target_simple_workbook()))
    info_ws = workbook["基本信息"]
    outcome_ws = workbook["课程目标"]
    info_ws["B4"] = "课程目标复用测试"
    info_ws["B3"] = ""
    for row in range(2, 5):
        outcome_ws.cell(row, 2).value = ""
        outcome_ws.cell(row, 3).value = ""
    stream = BytesIO()
    workbook.save(stream)
    workbook_bytes = stream.getvalue()

    db = SessionLocal()
    try:
        course = db.query(Course).filter(Course.course_code == "REUSE-COURSE-001").first()
        if not course:
            course = Course(
                course_code="REUSE-COURSE-001",
                course_name="课程目标复用测试",
                term="演示学期",
                department="计算机科学与技术系",
                major="示例专业",
                credit=3,
                owner="示例负责人",
                description="教学大纲导入课程",
            )
            db.add(course)
            db.flush()
        db.query(OBEOutcome).filter(OBEOutcome.course_id == course.id).delete()
        for index, description in enumerate(
            [
                "掌握示例课程的基本概念、过程和质量意识。",
                "能够完成示例问题获取、分析和模型表达。",
                "能够综合完成示例方案分析、评审与优化。",
            ],
            start=1,
        ):
            db.add(
                OBEOutcome(
                    course_id=course.id,
                    co_code=f"CO{index}",
                    co_name=f"课程目标{index}",
                    indicator=f"指标点{index}",
                    description=description,
                    threshold=0.7,
                )
            )
        db.commit()
    finally:
        db.close()

    with TestClient(app) as client:
        response = client.post(
            "/api/v1/import/teacher-workbook-task",
            files={
                "file": (
                    "可省略课程目标说明.xlsx",
                    workbook_bytes,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

        assert response.status_code == 200
        run_id = response.json()["data"]["run_overview"]["run"]["id"]
        dashboard_response = client.get(f"/api/v1/analysis-runs/{run_id}/dashboard")

    assert dashboard_response.status_code == 200
    dashboard = dashboard_response.json()["data"]
    outcomes = {item["co_code"]: item for item in dashboard["course_outcomes"]}
    assert "基本概念" in outcomes["CO1"]["description"]
    assert "示例问题获取" in outcomes["CO2"]["description"]
    assert "示例方案" in outcomes["CO3"]["description"]
    assert outcomes["CO1"]["indicator"] == "指标点1"
    assert outcomes["CO2"]["threshold"] == 0.7


def test_downloadable_teacher_input_template_is_simplified() -> None:
    workbook = load_workbook(TEMPLATE_PATH, data_only=True)

    assert workbook.sheetnames == ["填写说明", "基本信息", "试卷结构", "课程目标", "学生成绩"]
    assert "课程目标平均分和达成度" not in "".join(workbook.sheetnames)
    assert "课程考核结果分析图" not in "".join(workbook.sheetnames)

    guide_text = "\n".join(
        str(cell.value)
        for row in workbook["填写说明"].iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "平均分、分数段、达成度、图表和分析文字都由系统自动计算生成" in guide_text


def test_template_download_endpoint_serves_simplified_template_without_cache() -> None:
    with TestClient(app) as client:
        response = client.get("/api/v1/import/templates/teacher_input_template.xlsx")

    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store, max-age=0"
    assert quote("成绩输入模板_简洁版.xlsx") in response.headers["content-disposition"]

    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert workbook.sheetnames == ["填写说明", "基本信息", "试卷结构", "课程目标", "学生成绩"]
    assert workbook["填写说明"]["A1"].value == "成绩输入模板填写说明"


def test_teacher_workbook_report_endpoint_returns_template_shaped_docx() -> None:
    workbook_bytes = _build_teacher_workbook()
    with TestClient(app) as client:
        response = client.post(
            "/api/v1/import/teacher-workbook-report",
            files={
                "file": (
                    "老师模板.xlsx",
                    workbook_bytes,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert response.status_code == 200
    assert response.content[:2] == b"PK"

    document = Document(BytesIO(response.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "试 卷 分 析 表" in text
    assert "可视化补充页" not in text
    assert "可视化补充页" not in table_text
    assert "示例课程A" in table_text
    assert "1、试题对课程目标的支撑度分析" in table_text
    assert "2、学生作答情况对课程目标的达成度分析" in table_text
    assert "3、教师对今后教学持续改进的具体意见" in table_text
    assert "0、成绩统计摘要" not in table_text
    assert "课程目标1：能够围绕课程核心知识" in table_text
    assert "课程目标1：能够围绕课程核心知识" in table_text
    assert "填空题" in table_text
    assert "判断题" in table_text
    assert "计算题" in table_text
    assert "分析题" in table_text
    assert "课程目标达成度表" in table_text
    assert "任课教师签名" in table_text
    assert len(document.inline_shapes) >= 2
    assert document.tables[0].cell(8, 1).tables
    assert len(document.tables[0].cell(8, 1).tables[0].columns) == 8


def test_teacher_workbook_preview_endpoint_returns_clear_io_contract() -> None:
    workbook_bytes = _build_teacher_workbook()
    with TestClient(app) as client:
        response = client.post(
            "/api/v1/import/teacher-workbook-preview",
            files={
                "file": (
                    "老师模板.xlsx",
                    workbook_bytes,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"exam_date": "2025年6月22日"},
        )

    assert response.status_code == 200
    payload = response.json()["data"]
    assert payload["input_mode"] == "teacher_workbook"
    assert payload["meta"]["course_name"] == "示例课程A"
    assert payload["meta"]["exam_date"] == "2025年6月22日"
    assert payload["supported_outputs"]
    assert payload["pipeline"]
    assert payload["input_requirements"]
    assert payload["data_sources"]
    assert any(item["source_method"] == "computed" for item in payload["data_sources"])


def test_teacher_workbook_task_endpoint_creates_real_analysis_run() -> None:
    workbook_bytes = _build_teacher_workbook()
    with TestClient(app) as client:
        response = client.post(
            "/api/v1/import/teacher-workbook-task",
            files={
                "file": (
                    "老师模板.xlsx",
                    workbook_bytes,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"exam_date": "2025年6月22日"},
        )

        assert response.status_code == 200
        payload = response.json()["data"]
        assert payload["run_overview"]["course_name"] == "示例课程A"
        assert payload["run_overview"]["run"]["class_name"] == "示例班级A"
        assert payload["run_overview"]["run"]["status"] == "ready"
        assert payload["template_binding"]["input_template"] == "teacher_input_template.xlsx"

        list_response = client.get("/api/v1/analysis-runs")
        assert list_response.status_code == 200
        rows = list_response.json()["data"]
        assert rows
        imported_run_id = payload["run_overview"]["run"]["id"]
        assert any(
            item["run"]["id"] == imported_run_id and item["course_name"] == "示例课程A"
            for item in rows
        )

        dashboard_response = client.get(f"/api/v1/analysis-runs/{imported_run_id}/dashboard")
        assert dashboard_response.status_code == 200
        dashboard = dashboard_response.json()["data"]
        outcomes = {item["co_code"]: item for item in dashboard["course_outcomes"]}
        assert outcomes["CO2"]["supporting_groups"] == ["填空题"]
        assert outcomes["CO3"]["supporting_groups"] == ["判断题", "计算题", "分析题"]
        assert outcomes["CO3"]["full_score"] == 50
        assert outcomes["CO3"]["achievement"] == 0.8
        analysis_group = next(item for item in dashboard["question_groups"] if item["qgroup_name"] == "分析题")
        assert analysis_group["avg_score"] == 24
        assert analysis_group["achievement"] == 0.8


def test_teacher_workbook_task_removes_stale_orphan_students() -> None:
    db = SessionLocal()
    try:
        db.add(
            Student(
                student_no="90分以上学生人数：0",
                name="90分以上学生人数：0",
                class_name="示例班级A",
                major="示例专业",
                grade_year="2026",
            )
        )
        db.commit()
    finally:
        db.close()

    workbook_bytes = _build_teacher_workbook()
    with TestClient(app) as client:
        response = client.post(
            "/api/v1/import/teacher-workbook-task",
            files={
                "file": (
                    "老师模板.xlsx",
                    workbook_bytes,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"exam_date": "2025年6月22日"},
        )

    assert response.status_code == 200

    db = SessionLocal()
    try:
        stale_student = db.query(Student).filter(Student.student_no == "90分以上学生人数：0").first()
        class_students = db.query(Student).filter(Student.class_name == "示例班级A").all()
    finally:
        db.close()

    assert stale_student is None
    assert {"9000001", "9000002"}.issubset({student.student_no for student in class_students})
