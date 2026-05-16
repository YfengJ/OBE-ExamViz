from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from backend.app.models.course import Course
from backend.app.models.obe_outcome import OBEOutcome


def create_course_from_syllabus(db: Session, file: UploadFile) -> dict[str, Any]:
    content = file.file.read()
    file.file.seek(0)
    if not content:
        raise HTTPException(status_code=400, detail="教学大纲文件为空")

    text = _extract_text(content, file.filename or "")
    parsed = _parse_syllabus_text(text, file.filename or "")
    if not parsed["course_code"] or not parsed["course_name"]:
        raise HTTPException(status_code=400, detail="未能从教学大纲中识别课程代码和课程名称")

    course = db.query(Course).filter(Course.course_code == parsed["course_code"]).first()
    payload = {
        "course_code": parsed["course_code"],
        "course_name": parsed["course_name"],
        "term": parsed["term"] or "待补充",
        "department": parsed["department"] or "待补充",
        "major": parsed["major"] or "待补充",
        "credit": parsed["credit"],
        "owner": parsed["owner"] or None,
        "description": parsed["description"] or None,
    }
    if course:
        for key, value in payload.items():
            setattr(course, key, value)
    else:
        course = Course(**payload)
        db.add(course)
        db.flush()

    db.query(OBEOutcome).filter(OBEOutcome.course_id == course.id).delete()
    for index, outcome in enumerate(parsed["outcomes"], start=1):
        db.add(
            OBEOutcome(
                course_id=course.id,
                co_code=f"CO{index}",
                co_name=f"课程目标{index}",
                indicator=outcome.get("indicator") or None,
                description=outcome["description"],
                threshold=0.65,
            )
        )
    db.commit()
    db.refresh(course)
    outcomes = db.query(OBEOutcome).filter(OBEOutcome.course_id == course.id).order_by(OBEOutcome.co_code.asc()).all()
    return {"course": course, "outcomes": outcomes}


def _extract_text(content: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        return _extract_docx_text(content)
    if suffix in {".xlsx", ".xls", ".xlsm"}:
        return _extract_workbook_text(content)
    if suffix in {".txt", ".md"}:
        return content.decode("utf-8", errors="ignore")
    raise HTTPException(status_code=400, detail="暂支持 docx、xlsx/xls/xlsm 和 txt 格式的教学大纲")


def _extract_docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append("：".join(cells))
    return "\n".join(lines)


def _extract_workbook_text(content: bytes) -> str:
    workbook = load_workbook(BytesIO(content), data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(sheet.title)
        for row in sheet.iter_rows(values_only=True):
            cells = [str(value).strip() for value in row if value not in (None, "")]
            if cells:
                lines.append("：".join(cells))
    return "\n".join(lines)


def _parse_syllabus_text(text: str, filename: str) -> dict[str, Any]:
    normalized = re.sub(r"[ \t]+", " ", text.replace("\r\n", "\n").replace("\r", "\n"))
    return {
        "course_code": _extract_field(normalized, "课程代码", "课程编码", "课程编号") or _filename_code(filename),
        "course_name": _extract_field(normalized, "课程名称", "课程中文名称", "课程名"),
        "term": _extract_field(normalized, "开设学期", "开课学期", "开课时间", "建议修读学期"),
        "department": _extract_field(normalized, "开设学院", "开课单位", "承担单位", "院系", "开课学院"),
        "major": _extract_field(normalized, "适用专业", "授课对象", "面向专业"),
        "credit": _extract_credit(normalized),
        "owner": _extract_field(normalized, "课程负责人", "负责人", "任课教师", "主讲教师"),
        "description": _extract_description(normalized),
        "outcomes": _extract_outcomes(normalized),
    }


def _extract_field(text: str, *labels: str) -> str:
    label_pattern = "|".join(re.escape(label) for label in labels)
    patterns = [
        rf"(?:{label_pattern})\s*[:：]\s*([^\n]+)",
        rf"(?:{label_pattern})\s+([^\n]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return _clean_inline_value(match.group(1))
    return ""


def _extract_credit(text: str) -> float:
    match = re.search(r"学时\s*/\s*学分\s*[:：]\s*[\d.]+\s*/\s*([\d.]+)", text)
    if match:
        return float(match.group(1))
    direct = _extract_field(text, "学分", "课程学分")
    if "/" in direct:
        direct = direct.rsplit("/", 1)[-1]
    direct_number = _first_number(direct)
    if direct_number:
        return direct_number
    return 0.0


def _extract_description(text: str) -> str:
    match = re.search(r"(?:课程简介|课程描述|课程性质与任务)\s*[:：]?\s*(.*?)(?=\n\s*(?:[一二三四五六七八九十]、)?课程目标|\n\s*[一二三四五六七八九十]、|\Z)", text, flags=re.S)
    if match:
        return _clean_block(match.group(1))
    description = _extract_field(text, "课程简介", "课程描述")
    return description


def _extract_outcomes(text: str) -> list[dict[str, str]]:
    scoped_text = _course_goal_scope(text)
    rows: list[dict[str, str]] = []
    pattern = re.compile(
        r"课程目标\s*([0-9一二三四五六七八九十]+)\s*[:：]\s*(.*?)(?=\n\s*课程目标\s*[0-9一二三四五六七八九十]+\s*[:：]|\n\s*[四五六七八九十]、|\Z)",
        flags=re.S,
    )
    for match in pattern.finditer(scoped_text):
        body = _clean_block(match.group(2))
        if not body:
            continue
        indicator = ""
        indicator_match = re.search(r"(指标点\s*\d+(?:[-.]\d+)?)", body)
        if indicator_match:
            indicator = indicator_match.group(1).replace(" ", "")
        rows.append({"description": body, "indicator": indicator})
    return rows


def _course_goal_scope(text: str) -> str:
    start_markers = (
        r"\n\s*（一）\s*课程目标\s*\n",
        r"\n\s*\(一\)\s*课程目标\s*\n",
        r"\n\s*三、\s*课程目标\s*\n",
        r"\n\s*3[.、]\s*课程目标\s*\n",
    )
    start = 0
    for marker in start_markers:
        match = re.search(marker, text)
        if match:
            start = match.end()
            break
    scoped = text[start:]
    end_patterns = (
        r"\n\s*（二）\s*课程目标对",
        r"\n\s*\(二\)\s*课程目标对",
        r"\n\s*四、",
        r"\n\s*4[.、]",
    )
    end_positions = [match.start() for pattern in end_patterns if (match := re.search(pattern, scoped))]
    if end_positions:
        scoped = scoped[: min(end_positions)]
    return scoped


def _filename_code(filename: str) -> str:
    stem = Path(filename).stem.strip()
    safe = re.sub(r"[^A-Za-z0-9_-]+", "", stem)
    return safe[:24]


def _clean_inline_value(value: str) -> str:
    value = re.split(r"\s{2,}|；|;", value.strip())[0]
    return value.strip(" ：:\t")


def _clean_block(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip(" ：:")


def _first_number(value: str) -> float:
    match = re.search(r"\d+(?:\.\d+)?", value or "")
    return float(match.group(0)) if match else 0.0
