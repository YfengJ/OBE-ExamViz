from __future__ import annotations

import copy
import re
from collections import defaultdict
from io import BytesIO
from pathlib import Path
from typing import Any

import matplotlib
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager

from backend.app.ai.deepseek_client import DeepSeekClient
from backend.app.services.report_lineage import build_data_source_overview, build_input_requirements


ROOT_DIR = Path(__file__).resolve().parents[3]
TEMPLATE_DOCX_PATH = ROOT_DIR / "backend" / "templates" / "teacher_report_template.docx"
SIMPLE_INPUT_SHEETS = ("基本信息", "试卷结构", "课程目标", "学生成绩")
ACHIEVEMENT_INFO_SHEET = "课程基本信息"
ACHIEVEMENT_CALC_SHEETS = ("计算表", "达成度计算表")
ACHIEVEMENT_MAX_SCAN_COLUMNS = 200
DEFAULT_DEPARTMENT = "计算机科学与技术系"

COURSE_GOAL_DESCRIPTIONS: dict[str, list[dict[str, str]]] = {}


def load_teacher_workbook_context(file: UploadFile, exam_date_override: str | None = None) -> dict[str, Any]:
    content = file.file.read()
    file.file.seek(0)
    return load_teacher_workbook_context_from_bytes(
        content,
        filename=file.filename,
        exam_date_override=exam_date_override,
    )


def load_teacher_workbook_context_from_bytes(
    content: bytes,
    filename: str | None = None,
    exam_date_override: str | None = None,
) -> dict[str, Any]:
    if not content:
        raise HTTPException(status_code=400, detail="上传文件为空")

    workbook_formula, workbook_value = _load_excel_workbooks(content, filename=filename)
    context = _parse_teacher_workbook(workbook_formula, workbook_value, exam_date_override=exam_date_override)
    return {
        "source_filename": filename or "",
        "recognized_sheets": workbook_value.sheetnames,
        "context": context,
    }


def _load_excel_workbooks(content: bytes, filename: str | None = None):
    try:
        return (
            load_workbook(BytesIO(content), data_only=False),
            load_workbook(BytesIO(content), data_only=True),
        )
    except Exception as openpyxl_error:
        if not _looks_like_legacy_xls(content, filename):
            raise HTTPException(status_code=400, detail=f"无法读取 Excel 工作簿：{openpyxl_error}") from openpyxl_error
        try:
            legacy_workbook = _load_legacy_xls_as_openpyxl(content)
        except HTTPException:
            raise
        except Exception as legacy_error:
            raise HTTPException(status_code=400, detail=f"无法读取旧版 .xls 工作簿：{legacy_error}") from legacy_error
        return legacy_workbook, legacy_workbook


def _looks_like_legacy_xls(content: bytes, filename: str | None = None) -> bool:
    lower_name = (filename or "").lower()
    return lower_name.endswith(".xls") or content.startswith(b"\xd0\xcf\x11\xe0")


def _load_legacy_xls_as_openpyxl(content: bytes):
    try:
        import xlrd
    except ImportError as exc:
        raise HTTPException(status_code=400, detail="当前环境缺少 xlrd，无法读取旧版 .xls 成绩文件") from exc

    legacy = xlrd.open_workbook(file_contents=content)
    workbook = Workbook()
    default_sheet = workbook.active
    used_titles: set[str] = set()
    for sheet_index, legacy_sheet in enumerate(legacy.sheets()):
        sheet = default_sheet if sheet_index == 0 else workbook.create_sheet()
        sheet.title = _unique_sheet_title(legacy_sheet.name, used_titles)
        used_titles.add(sheet.title)
        for row_index in range(legacy_sheet.nrows):
            for column_index in range(legacy_sheet.ncols):
                value = _xlrd_cell_value(legacy_sheet.cell(row_index, column_index), legacy.datemode)
                if value not in (None, ""):
                    sheet.cell(row=row_index + 1, column=column_index + 1).value = value
    return workbook


def _unique_sheet_title(title: str, used_titles: set[str]) -> str:
    base = re.sub(r"[:\\/?*\[\]]", "_", (title or "Sheet").strip())[:31] or "Sheet"
    candidate = base
    suffix = 1
    while candidate in used_titles:
        tail = f"_{suffix}"
        candidate = f"{base[:31 - len(tail)]}{tail}"
        suffix += 1
    return candidate


def _xlrd_cell_value(cell, datemode: int) -> Any:
    try:
        import xlrd
    except ImportError:
        return cell.value
    if cell.ctype == xlrd.XL_CELL_EMPTY:
        return ""
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate.xldate_as_datetime(cell.value, datemode)
        except Exception:
            return cell.value
    if cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
        return int(cell.value)
    return cell.value


def preview_teacher_template_workbook(file: UploadFile, exam_date_override: str | None = None) -> dict[str, Any]:
    parsed = load_teacher_workbook_context(file, exam_date_override=exam_date_override)
    return _preview_payload(parsed["source_filename"], parsed["recognized_sheets"], parsed["context"])


def build_teacher_template_report(file: UploadFile, exam_date_override: str | None = None) -> tuple[bytes, str]:
    parsed = load_teacher_workbook_context(file, exam_date_override=exam_date_override)
    context = parsed["context"]
    document = _build_report_document([context])
    return _document_to_bytes(document), _suggest_filename(parsed["source_filename"], context)


async def build_teacher_template_report_with_ai(file: UploadFile, exam_date_override: str | None = None) -> tuple[bytes, str]:
    parsed = load_teacher_workbook_context(file, exam_date_override=exam_date_override)
    context = await enrich_context_with_ai_narrative(parsed["context"])
    document = _build_report_document([context])
    return _document_to_bytes(document), _suggest_filename(parsed["source_filename"], context)


def build_teacher_template_report_from_context(context: dict[str, Any], filename: str | None = None) -> tuple[bytes, str]:
    document = _build_report_document([context])
    return _document_to_bytes(document), _suggest_filename(filename, context)


def build_teacher_template_report_from_analysis_context(context: dict[str, Any]) -> bytes:
    document = _build_report_document([_normalize_analysis_context(context)])
    return _document_to_bytes(document)


async def build_teacher_template_report_from_analysis_context_with_ai(context: dict[str, Any]) -> bytes:
    normalized = _normalize_analysis_context(context)
    enriched = await enrich_context_with_ai_narrative(normalized)
    document = _build_report_document([enriched])
    return _document_to_bytes(document)


def _preview_payload(filename: str | None, sheet_names: list[str], context: dict[str, Any]) -> dict[str, Any]:
    ai_client = DeepSeekClient()
    return {
        "input_mode": "teacher_workbook",
        "source_filename": filename or "",
        "recognized_sheets": sheet_names,
        "meta": context["meta"],
        "score_stats": context["score_stats"],
        "difficulty_label": context["difficulty_label"],
        "question_group_count": len(context["question_groups"]),
        "question_item_count": len(context["question_items"]),
        "course_outcome_count": len(context["course_outcomes"]),
        "warning_count": len(context["warnings"]),
        "ai_enabled": ai_client.is_configured,
        "narrative_engine": "deepseek" if ai_client.is_configured else "rule_based",
        "supported_outputs": [
            {
                "key": "teacher_docx",
                "label": "正式 Word 报告",
                "description": "按正式版式回填基础信息、成绩统计和教学质量分析文字。",
            },
            {
                "key": "structured_analysis",
                "label": "分析结果",
                "description": "生成成绩分布、课程目标达成度和题型得分率。",
            },
        ],
        "input_requirements": build_input_requirements(context),
        "data_sources": build_data_source_overview(context),
        "pipeline": [
            "读取 Excel 工作簿",
            "计算成绩统计、题型表现和课程目标达成度",
            "生成试卷分析表正文",
            "导出 Word 报告",
        ],
    }


def _suggest_filename(filename: str | None, context: dict[str, Any]) -> str:
    stem = Path(filename or "teacher_workbook").stem
    class_name = str(context["meta"].get("class_name") or "班级")
    return f"{stem}_{class_name}_试卷分析报告.docx"


def _parse_teacher_workbook(workbook_formula, workbook_value, exam_date_override: str | None = None) -> dict[str, Any]:
    if _has_simple_input_sheets(workbook_value.sheetnames):
        return _parse_simple_input_workbook(workbook_value, exam_date_override=exam_date_override)
    if _has_achievement_analysis_sheets(workbook_value.sheetnames):
        return _parse_achievement_analysis_workbook(workbook_value, exam_date_override=exam_date_override)

    raw_sheet_name = _find_sheet_name(workbook_value.sheetnames, include=("期末",), exclude=("分析", "目标"))
    main_sheet_name = _find_sheet_name(workbook_value.sheetnames, include=("班",), exclude=("目标", "分析图", "平均分", "期末"))
    summary_sheet_name = _find_sheet_name(workbook_value.sheetnames, include=("课程目标平均分和达成度",))
    chart_sheet_name = _find_sheet_name(workbook_value.sheetnames, include=("课程考核结果分析图",))

    if not raw_sheet_name or not main_sheet_name or not summary_sheet_name:
        score_only_context = _parse_score_only_workbook(workbook_value, exam_date_override=exam_date_override)
        if score_only_context:
            return score_only_context
        raise HTTPException(status_code=400, detail="上传的 Excel 不符合成绩工作簿结构")

    raw_ws = workbook_value[raw_sheet_name]
    main_ws_value = workbook_value[main_sheet_name]
    main_ws_formula = workbook_formula[main_sheet_name]
    summary_ws = workbook_value[summary_sheet_name]
    chart_ws = workbook_value[chart_sheet_name] if chart_sheet_name else None

    meta = _parse_meta(raw_ws, main_ws_value, exam_date_override=exam_date_override)
    students, question_items = _parse_students_and_questions(raw_ws, main_ws_value)
    score_stats = _score_stats(students)
    score_segments = _score_segments(students)
    question_groups = _question_group_summary(students, question_items)
    threshold_map = _threshold_map(chart_ws)
    support_map = _supporting_groups(main_ws_formula, main_ws_value)
    course_outcomes = _course_outcomes(summary_ws, threshold_map, support_map)
    student_outcomes = _student_outcomes(main_ws_value, threshold_map)
    overall_achievement = _overall_achievement(student_outcomes)
    warnings = _warning_rows(students, overall_achievement)
    narrative_sections = _build_teacher_narrative_sections(
        meta=meta,
        score_stats=score_stats,
        score_segments=score_segments,
        question_groups=question_groups,
        course_outcomes=course_outcomes,
        student_outcomes=student_outcomes,
        overall_achievement=overall_achievement,
    )

    return {
        "meta": meta,
        "students": students,
        "score_stats": score_stats,
        "score_segments": score_segments,
        "difficulty_label": _difficulty_label(score_stats["average_score"], 100.0),
        "question_items": question_items,
        "question_groups": question_groups,
        "course_outcomes": course_outcomes,
        "student_outcomes": student_outcomes,
        "overall_achievement": overall_achievement,
        "warnings": warnings,
        "narrative_sections": narrative_sections,
        "narrative": _combine_narrative_sections(narrative_sections),
    }


def _find_sheet_name(sheet_names: list[str], include: tuple[str, ...], exclude: tuple[str, ...] = ()) -> str | None:
    for name in sheet_names:
        if all(token in name for token in include) and not any(token in name for token in exclude):
            return name
    return None


def _has_simple_input_sheets(sheet_names: list[str]) -> bool:
    normalized = {name.strip() for name in sheet_names}
    return all(name in normalized for name in SIMPLE_INPUT_SHEETS)


def _has_achievement_analysis_sheets(sheet_names: list[str]) -> bool:
    normalized = {name.strip() for name in sheet_names}
    return ACHIEVEMENT_INFO_SHEET in normalized and any(name in normalized for name in ACHIEVEMENT_CALC_SHEETS)


def _get_sheet_by_normalized_name(workbook, expected_name: str):
    for name in workbook.sheetnames:
        if name.strip() == expected_name:
            return workbook[name]
    raise HTTPException(status_code=400, detail=f"缺少工作表：{expected_name}")


def _get_first_sheet_by_normalized_name(workbook, *expected_names: str):
    expected = {name.strip() for name in expected_names}
    for name in workbook.sheetnames:
        if name.strip() in expected:
            return workbook[name]
    raise HTTPException(status_code=400, detail=f"缺少工作表：{' / '.join(expected_names)}")


def _parse_achievement_analysis_workbook(workbook_value, exam_date_override: str | None = None) -> dict[str, Any]:
    info_ws = _get_sheet_by_normalized_name(workbook_value, "课程基本信息")
    calc_ws = _get_first_sheet_by_normalized_name(workbook_value, *ACHIEVEMENT_CALC_SHEETS)

    meta = _parse_achievement_meta(info_ws, exam_date_override=exam_date_override)
    students, question_items = _parse_achievement_students_and_questions(calc_ws)
    if not students:
        raise HTTPException(status_code=400, detail="计算表至少需要包含一名学生的成绩")
    if not question_items:
        raise HTTPException(status_code=400, detail="计算表未识别到题型分值和课程目标对应关系")

    threshold = float(meta.get("outcome_threshold") or 0.65)
    course_outcomes = _course_outcomes_from_question_items(
        question_items,
        students,
        default_threshold=threshold,
        outcome_inputs=_achievement_outcome_inputs(info_ws, threshold),
    )
    meta["student_count_actual"] = len(students)
    if not meta.get("student_count_expected"):
        meta["student_count_expected"] = len(students)

    score_stats = _score_stats(students)
    score_segments = _score_segments(students)
    question_groups = _question_group_summary(students, question_items)
    student_outcomes = _simple_student_outcomes(students, question_items, course_outcomes)
    overall_achievement = _overall_achievement(student_outcomes)
    warnings = _warning_rows(students, overall_achievement)
    narrative_sections = _build_teacher_narrative_sections(
        meta=meta,
        score_stats=score_stats,
        score_segments=score_segments,
        question_groups=question_groups,
        course_outcomes=course_outcomes,
        student_outcomes=student_outcomes,
        overall_achievement=overall_achievement,
    )

    return {
        "meta": meta,
        "students": students,
        "score_stats": score_stats,
        "score_segments": score_segments,
        "difficulty_label": _difficulty_label(score_stats["average_score"], _simple_exam_total(question_items)),
        "question_items": question_items,
        "question_groups": question_groups,
        "course_outcomes": course_outcomes,
        "student_outcomes": student_outcomes,
        "overall_achievement": overall_achievement,
        "warnings": warnings,
        "narrative_sections": narrative_sections,
        "narrative": _combine_narrative_sections(narrative_sections),
    }


def _parse_achievement_meta(ws, exam_date_override: str | None = None) -> dict[str, Any]:
    course_code = _sheet_value_by_labels(ws, "课程代码", "课程编码", "课程编号")
    course_name = _sheet_value_by_labels(ws, "课程名称", "课程名")
    class_name = _sheet_value_by_labels(ws, "授课对象", "教学班级", "班级", "专业班级")
    expected = _safe_float(_sheet_value_by_labels(ws, "学生人数", "应考人数"))
    threshold = _safe_float(_sheet_value_by_labels(ws, "期望达成标准", "达成标准", "课程目标达成阈值"))
    if threshold > 1:
        threshold = threshold / 100
    return {
        "academic_year": _sheet_value_by_labels(ws, "学年学期", "学期", "开设学期"),
        "department": _sheet_value_by_labels(ws, "开设学院", "承担单位", "开课单位", "院系"),
        "course_code": course_code,
        "course_name": course_name,
        "teacher_name": _sheet_value_by_labels(ws, "任课教师", "教师", "课程负责人"),
        "class_name": class_name,
        "student_count_expected": int(expected) if expected else 0,
        "student_count_actual": 0,
        "exam_date": (exam_date_override or "").strip() or _sheet_value_by_labels(ws, "考试日期", "考试时间"),
        "outcome_threshold": threshold or 0.65,
    }


def _sheet_value_by_labels(ws, *labels: str) -> str:
    label_set = {_normalize_header(label) for label in labels}
    for row in range(1, ws.max_row + 1):
        for column in range(1, ws.max_column + 1):
            value = _normalize_header(ws.cell(row, column).value)
            if value not in label_set:
                continue
            for next_column in range(column + 1, min(ws.max_column, column + 4) + 1):
                candidate = _text(ws.cell(row, next_column).value)
                if candidate:
                    return candidate
            for next_row in range(row + 1, min(ws.max_row, row + 3) + 1):
                candidate = _text(ws.cell(next_row, column).value)
                if candidate:
                    return candidate
    return ""


def _parse_achievement_students_and_questions(ws) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    scan_max_column = min(ws.max_column, ACHIEVEMENT_MAX_SCAN_COLUMNS)
    header_row = _find_achievement_header_row(ws, scan_max_column)
    if not header_row:
        raise HTTPException(status_code=400, detail="计算表需要包含学号、姓名和题型得分列")

    header = _header_map_limited(ws, header_row, scan_max_column)
    student_no_col = header.get("学号")
    name_col = header.get("姓名")
    class_col = header.get("班级")
    total_col = next((header.get(name) for name in ("期末总分", "总分", "成绩") if header.get(name)), None)
    data_start_row = _achievement_data_start_row(ws, header_row, student_no_col)
    full_score_row = _achievement_full_score_row(ws, header_row, data_start_row, scan_max_column)
    target_row = max(1, header_row - 1)
    group_to_co = _achievement_group_to_co(ws, scan_max_column)

    used_labels: set[str] = set()
    question_columns: list[dict[str, Any]] = []
    question_items: list[dict[str, Any]] = []
    skipped_headers = {"序号", "学号", "姓名", "班级", "期末总分", "总分", "成绩"}
    for column in range(1, scan_max_column + 1):
        qtype = _text(ws.cell(header_row, column).value)
        normalized = _normalize_header(qtype)
        if not qtype or normalized in skipped_headers or "达成度" in normalized:
            continue
        full_score = _safe_float(ws.cell(full_score_row, column).value)
        if full_score <= 0:
            full_score = _safe_float(ws.cell(header_row + 1, column).value)
        if full_score <= 0:
            continue
        co_code = _achievement_co_for_column(ws, target_row, column) or group_to_co.get(qtype) or f"CO{len(question_items) + 1}"
        label = _simple_question_label(qtype, "", str(len(question_items) + 1), used_labels)
        used_labels.add(label)
        question_columns.append({"label": label, "column": column})
        question_items.append(
            {
                "label": label,
                "qno": str(len(question_items) + 1),
                "qgroup_name": qtype,
                "full_score": round(full_score, 2),
                "co_code": co_code,
                "avg_score": 0.0,
                "difficulty": 0.0,
                "pass_rate": 0.0,
            }
        )

    students: list[dict[str, Any]] = []
    empty_streak = 0
    for row in range(data_start_row, ws.max_row + 1):
        row_marker = " ".join(_text(ws.cell(row, column).value) for column in range(1, min(ws.max_column, 8) + 1))
        if "合计" in row_marker or "平均" in row_marker:
            break
        student_no = _normalize_student_no(ws.cell(row, student_no_col).value) if student_no_col else ""
        if not student_no:
            empty_streak += 1
            if empty_streak >= 5:
                break
            continue
        empty_streak = 0
        question_scores = {
            item["label"]: _safe_float(ws.cell(row, item["column"]).value)
            for item in question_columns
        }
        computed_total = round(sum(question_scores.values()), 2)
        exam_total = _simple_exam_total(question_items)
        if total_col:
            final_score = _safe_float(ws.cell(row, total_col).value)
        elif exam_total and exam_total != 100:
            final_score = int((computed_total / exam_total * 100) + 0.5)
        else:
            final_score = computed_total
        students.append(
            {
                "student_no": student_no,
                "name": _text(ws.cell(row, name_col).value) if name_col else "",
                "class_name": _text(ws.cell(row, class_col).value) if class_col else "",
                "question_scores": question_scores,
                "final_score": round(final_score or computed_total, 2),
            }
        )

    _attach_question_statistics(question_items, students)
    return students, question_items


def _find_achievement_header_row(ws, max_column: int) -> int | None:
    for row in range(1, min(ws.max_row, 30) + 1):
        headers = {_normalize_header(ws.cell(row, column).value) for column in range(1, max_column + 1)}
        if "学号" in headers and "姓名" in headers:
            return row
    return None


def _header_map_limited(ws, row: int, max_column: int) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for column in range(1, max_column + 1):
        name = _normalize_header(ws.cell(row, column).value)
        if name:
            mapping[name] = column
    return mapping


def _achievement_data_start_row(ws, header_row: int, student_no_col: int | None) -> int:
    if not student_no_col:
        return header_row + 1
    for row in range(header_row + 1, ws.max_row + 1):
        student_no = _normalize_student_no(ws.cell(row, student_no_col).value)
        if student_no:
            return row
    return header_row + 1


def _achievement_full_score_row(ws, header_row: int, data_start_row: int, max_column: int) -> int:
    best_row = header_row + 1
    best_count = -1
    for row in range(header_row + 1, max(header_row + 2, data_start_row)):
        count = sum(1 for column in range(1, max_column + 1) if _safe_float(ws.cell(row, column).value) > 0)
        if count > best_count:
            best_count = count
            best_row = row
    return best_row


def _achievement_co_for_column(ws, target_row: int, column: int) -> str:
    for current_column in range(column, 0, -1):
        value = _normalize_co_code(_resolve_merged_value(ws, target_row, current_column) or ws.cell(target_row, current_column).value)
        if re.fullmatch(r"CO\d+", value):
            return value
    return ""


def _achievement_group_to_co(ws, max_column: int) -> dict[str, str]:
    mapping: dict[str, str] = {}
    target_cells: list[tuple[int, int, str]] = []
    for row in range(1, min(ws.max_row, 40) + 1):
        for column in range(1, max_column + 1):
            raw_value = ws.cell(row, column).value
            if not _looks_like_course_target_label(raw_value):
                continue
            co_code = _normalize_co_code(raw_value)
            if re.fullmatch(r"CO\d+", co_code):
                target_cells.append((row, column, co_code))

    for index, (row, column, co_code) in enumerate(target_cells):
        next_row = target_cells[index + 1][0] if index + 1 < len(target_cells) and target_cells[index + 1][1] == column else ws.max_row + 1
        for current_row in range(row + 1, min(next_row, row + 8)):
            for current_column in range(column, min(max_column, column + 6) + 1):
                group_name = _text(ws.cell(current_row, current_column).value)
                if group_name and not _normalize_co_code(group_name).startswith("CO"):
                    mapping[group_name] = co_code
    return mapping


def _achievement_outcome_inputs(ws, default_threshold: float) -> dict[str, dict[str, Any]]:
    rows: dict[str, dict[str, Any]] = {}
    for row in range(1, ws.max_row + 1):
        for column in range(1, ws.max_column + 1):
            raw_value = ws.cell(row, column).value
            if not _looks_like_course_target_label(raw_value):
                continue
            co_code = _normalize_co_code(raw_value)
            if not re.fullmatch(r"CO\d+", co_code):
                continue
            description = _text(ws.cell(row, column + 1).value) if column < ws.max_column else ""
            rows.setdefault(
                co_code,
                {
                    "threshold": default_threshold,
                    "indicator": "",
                    "description": description if "课程目标" not in description else "",
                },
            )
    return rows


def _looks_like_course_target_label(value: Any) -> bool:
    text = _normalize_header(value).upper()
    return bool(text.startswith("CO") or "课程目标" in text)


def _course_outcomes_from_question_items(
    question_items: list[dict[str, Any]],
    students: list[dict[str, Any]],
    *,
    default_threshold: float = 0.65,
    outcome_inputs: dict[str, dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    outcome_inputs = outcome_inputs or {}
    co_codes = sorted({item["co_code"] for item in question_items}, key=_co_sort_key)
    total_full = _simple_exam_total(question_items) or 1.0
    rows: list[dict[str, Any]] = []
    for index, co_code in enumerate(co_codes, start=1):
        linked_items = [item for item in question_items if item["co_code"] == co_code]
        full_score = round(sum(float(item["full_score"]) for item in linked_items), 2)
        scores = [
            sum(float(student["question_scores"].get(item["label"], 0)) for item in linked_items)
            for student in students
        ]
        avg_score = round(sum(scores) / len(scores), 2) if scores else 0.0
        achievement = round(avg_score / full_score, 4) if full_score else 0.0
        input_row = outcome_inputs.get(co_code, {})
        label = f"课程目标{_co_number(co_code) or index}"
        supporting_groups: list[str] = []
        for item in linked_items:
            group_name = str(item["qgroup_name"])
            if group_name not in supporting_groups:
                supporting_groups.append(group_name)
        weight = round(full_score / total_full, 4)
        rows.append(
            {
                "label": label,
                "co_code": co_code,
                "co_name": label,
                "indicator": input_row.get("indicator", ""),
                "description": input_row.get("description", ""),
                "full_score": full_score,
                "avg_score": avg_score,
                "achievement": achievement,
                "threshold": round(float(input_row.get("threshold") or default_threshold), 4),
                "supporting_groups": supporting_groups,
                "weight": weight,
                "result": round(weight * achievement, 4),
            }
        )
    return rows


def _parse_score_only_workbook(workbook_value, exam_date_override: str | None = None) -> dict[str, Any] | None:
    candidates: list[dict[str, Any]] = []
    for sheet_name in workbook_value.sheetnames:
        sheet = workbook_value[sheet_name]
        parsed = _parse_score_only_sheet(sheet, exam_date_override=exam_date_override)
        if parsed:
            parsed["sheet_name"] = sheet_name
            candidates.append(parsed)
    if not candidates:
        return None

    selected = max(candidates, key=lambda item: len(item["students"]))
    meta = selected["meta"]
    students = selected["students"]
    question_items = selected["question_items"]
    _attach_question_statistics(question_items, students)
    course_outcomes = _course_outcomes_from_question_items(
        question_items,
        students,
        default_threshold=float(meta.get("outcome_threshold") or 0.65),
    )
    return _build_context_from_parts(
        meta=meta,
        students=students,
        question_items=question_items,
        course_outcomes=course_outcomes,
        input_mode="score_only_workbook",
    )


def _parse_score_only_sheet(ws, exam_date_override: str | None = None) -> dict[str, Any] | None:
    header_row = _find_score_only_header_row(ws)
    if not header_row:
        return None

    header = _header_map(ws, header_row)
    student_no_col = header.get("学号")
    if not student_no_col:
        return None
    name_col = header.get("姓名")
    class_col = header.get("班级") or header.get("教学班级") or header.get("专业班级")
    total_col = _score_only_total_column(header)
    data_start_row = _score_only_data_start_row(ws, header_row, student_no_col)
    if not data_start_row:
        return None

    question_columns = _score_only_question_columns(ws, header_row, data_start_row, total_col)
    students = _score_only_students(ws, data_start_row, student_no_col, name_col, class_col, total_col, question_columns)
    if not students:
        return None

    question_items = _score_only_question_items(ws, header_row, question_columns, students)
    if not question_items:
        question_items = _score_only_total_question(students)

    meta = _parse_score_only_meta(ws, header_row, exam_date_override=exam_date_override)
    if not meta.get("class_name"):
        meta["class_name"] = str(ws.title or "").strip()
    if not meta.get("student_count_expected"):
        meta["student_count_expected"] = len(students)
    meta["student_count_actual"] = len(students)

    return {"meta": meta, "students": students, "question_items": question_items}


def _find_score_only_header_row(ws) -> int | None:
    max_column = min(ws.max_column, 80)
    for row in range(1, min(ws.max_row, 40) + 1):
        headers = {_normalize_header(ws.cell(row, column).value) for column in range(1, max_column + 1)}
        if "学号" in headers and _score_only_total_column({header: 1 for header in headers}):
            return row
    return None


def _score_only_total_column(header: dict[str, int]) -> int | None:
    for name in ("期末总分", "卷面总分", "总分", "成绩", "期末成绩", "卷面成绩"):
        if name in header:
            return header[name]
    for name, column in header.items():
        if ("总分" in name or "成绩" in name) and "平时" not in name and "姓名" not in name:
            return column
    return None


def _score_only_data_start_row(ws, header_row: int, student_no_col: int) -> int | None:
    for row in range(header_row + 1, ws.max_row + 1):
        student_no = _normalize_student_no(ws.cell(row, student_no_col).value)
        if student_no:
            return row
    return None


def _score_only_question_columns(ws, header_row: int, data_start_row: int, total_col: int | None) -> list[dict[str, Any]]:
    skipped_headers = {
        "序号",
        "学号",
        "姓名",
        "班级",
        "教学班级",
        "专业班级",
        "总分",
        "成绩",
        "期末总分",
        "期末成绩",
        "卷面总分",
        "卷面成绩",
    }
    columns: list[dict[str, Any]] = []
    used_labels: set[str] = set()
    for column in range(1, min(ws.max_column, 80) + 1):
        if column == total_col:
            continue
        raw_header = _text(ws.cell(header_row, column).value)
        normalized = _normalize_header(raw_header)
        if not raw_header or normalized in skipped_headers:
            continue
        if not _score_column_has_numeric_values(ws, data_start_row, column):
            continue
        qtype = _score_only_question_name(raw_header)
        full_score = _full_score_from_header(raw_header)
        label = _simple_question_label(qtype, "", str(len(columns) + 1), used_labels)
        used_labels.add(label)
        columns.append(
            {
                "column": column,
                "label": label,
                "qgroup_name": qtype,
                "full_score": full_score,
            }
        )
    return columns


def _score_column_has_numeric_values(ws, data_start_row: int, column: int) -> bool:
    checked = 0
    numeric = 0
    empty_streak = 0
    for row in range(data_start_row, min(ws.max_row, data_start_row + 80) + 1):
        value = ws.cell(row, column).value
        if value in (None, ""):
            empty_streak += 1
            if empty_streak >= 5:
                break
            continue
        empty_streak = 0
        checked += 1
        if _safe_float(value) or str(value).strip() in {"0", "0.0"}:
            numeric += 1
    return checked > 0 and numeric >= max(1, checked // 2)


def _score_only_students(
    ws,
    data_start_row: int,
    student_no_col: int,
    name_col: int | None,
    class_col: int | None,
    total_col: int | None,
    question_columns: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    students: list[dict[str, Any]] = []
    empty_streak = 0
    for row in range(data_start_row, ws.max_row + 1):
        row_marker = " ".join(_text(ws.cell(row, column).value) for column in range(1, min(ws.max_column, 10) + 1))
        if "合计" in row_marker or "平均" in row_marker:
            break
        student_no = _normalize_student_no(ws.cell(row, student_no_col).value)
        if not student_no:
            empty_streak += 1
            if empty_streak >= 5:
                break
            continue
        empty_streak = 0
        question_scores = {
            item["label"]: _safe_float(ws.cell(row, item["column"]).value)
            for item in question_columns
        }
        computed_total = round(sum(question_scores.values()), 2)
        final_score = _safe_float(ws.cell(row, total_col).value) if total_col else computed_total
        students.append(
            {
                "student_no": student_no,
                "name": _text(ws.cell(row, name_col).value) if name_col else "",
                "class_name": _text(ws.cell(row, class_col).value) if class_col else "",
                "question_scores": question_scores,
                "final_score": round(final_score or computed_total, 2),
            }
        )
    return students


def _score_only_question_items(
    ws,
    header_row: int,
    question_columns: list[dict[str, Any]],
    students: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for index, column_info in enumerate(question_columns, start=1):
        scores = [float(student["question_scores"].get(column_info["label"], 0)) for student in students]
        full_score = float(column_info.get("full_score") or 0)
        if full_score <= 0 and scores:
            full_score = max(scores)
        if full_score <= 0:
            continue
        items.append(
            {
                "label": column_info["label"],
                "qno": str(index),
                "qgroup_name": column_info["qgroup_name"],
                "full_score": round(full_score, 2),
                "co_code": f"CO{index}",
                "avg_score": 0.0,
                "difficulty": 0.0,
                "pass_rate": 0.0,
            }
        )
    return items


def _score_only_total_question(students: list[dict[str, Any]]) -> list[dict[str, Any]]:
    for student in students:
        student["question_scores"] = {"期末总分": float(student.get("final_score") or 0)}
    max_score = max((float(student.get("final_score") or 0) for student in students), default=100.0)
    full_score = 100.0 if max_score <= 100 else max_score
    return [
        {
            "label": "期末总分",
            "qno": "1",
            "qgroup_name": "期末总分",
            "full_score": full_score,
            "co_code": "CO1",
            "avg_score": 0.0,
            "difficulty": 0.0,
            "pass_rate": 0.0,
        }
    ]


def _parse_score_only_meta(ws, header_row: int, exam_date_override: str | None = None) -> dict[str, Any]:
    leading_text = "\n".join(
        _text(ws.cell(row, column).value)
        for row in range(1, header_row)
        for column in range(1, min(ws.max_column, 12) + 1)
        if _text(ws.cell(row, column).value)
    )
    academic_year = _extract_academic_year_from_text(leading_text)
    expected = (
        _extract_number(r"(?:学生人数|上课人数|应考人数)[:：]\s*(\d+)", leading_text)
        or _extract_number(r"(?:学生人数|上课人数|应考人数)\s*(\d+)", leading_text)
        or 0
    )
    return {
        "academic_year": academic_year,
        "department": "",
        "course_code": "",
        "course_name": _extract(r"课程名称[:：]\s*([^\n\r]+)", leading_text),
        "teacher_name": _extract(r"(?:任课教师|代课教师|教师)[:：]\s*([^\n\r]+)", leading_text),
        "class_name": _extract(r"(?:专业班级|教学班级|班级)[:：]\s*([^\n\r]+)", leading_text),
        "student_count_expected": int(expected) if expected else 0,
        "student_count_actual": 0,
        "exam_date": (exam_date_override or "").strip(),
        "outcome_threshold": 0.65,
    }


def _extract_academic_year_from_text(text: str) -> str:
    match = re.search(r"(20\d{2}\s*[-—]\s*20\d{2}\s*学年第?[一二三四五六七八九十\d]+学期)", text)
    if not match:
        return ""
    value = re.sub(r"\s+", "", match.group(1)).replace("-", "—")
    return value


def _score_only_question_name(raw_header: str) -> str:
    text = re.sub(r"\s+", "", raw_header)
    text = re.sub(r"\d+(?:\.\d+)?分", "", text)
    return text or raw_header.strip() or "题目"


def _full_score_from_header(raw_header: str) -> float:
    match = re.search(r"(\d+(?:\.\d+)?)\s*分", str(raw_header or ""))
    return _safe_float(match.group(1)) if match else 0.0


def _normalize_student_no(value: Any) -> str:
    text = _text(value)
    if text.endswith(".0"):
        text = text[:-2]
    if not re.fullmatch(r"\d{1,20}", text):
        return ""
    return text


def _build_context_from_parts(
    *,
    meta: dict[str, Any],
    students: list[dict[str, Any]],
    question_items: list[dict[str, Any]],
    course_outcomes: list[dict[str, Any]],
    input_mode: str | None = None,
) -> dict[str, Any]:
    meta["student_count_actual"] = len(students)
    if not meta.get("student_count_expected"):
        meta["student_count_expected"] = len(students)
    score_stats = _score_stats(students)
    score_segments = _score_segments(students)
    question_groups = _question_group_summary(students, question_items)
    student_outcomes = _simple_student_outcomes(students, question_items, course_outcomes)
    overall_achievement = _overall_achievement(student_outcomes)
    warnings = _warning_rows(students, overall_achievement)
    narrative_sections = _build_teacher_narrative_sections(
        meta=meta,
        score_stats=score_stats,
        score_segments=score_segments,
        question_groups=question_groups,
        course_outcomes=course_outcomes,
        student_outcomes=student_outcomes,
        overall_achievement=overall_achievement,
    )
    context = {
        "meta": meta,
        "students": students,
        "score_stats": score_stats,
        "score_segments": score_segments,
        "difficulty_label": _difficulty_label(score_stats["average_score"], _simple_exam_total(question_items)),
        "question_items": question_items,
        "question_groups": question_groups,
        "course_outcomes": course_outcomes,
        "student_outcomes": student_outcomes,
        "overall_achievement": overall_achievement,
        "warnings": warnings,
        "narrative_sections": narrative_sections,
        "narrative": _combine_narrative_sections(narrative_sections),
    }
    if input_mode:
        context["input_mode"] = input_mode
    return context


def enrich_context_with_course_catalog(
    context: dict[str, Any],
    catalog_outcomes: list[dict[str, Any]],
    *,
    remap_score_only: bool = False,
) -> dict[str, Any]:
    if not catalog_outcomes:
        return context
    students = context.get("students") or []
    question_items = context.get("question_items") or []
    if not students:
        return context

    catalog = _normalize_catalog_outcomes(catalog_outcomes)
    if not catalog:
        return context

    is_score_only = context.get("input_mode") == "score_only_workbook"
    if is_score_only and len(question_items) < len(catalog):
        question_items = _synthesize_questions_from_final_scores(students, catalog, question_items)
        context["question_items"] = question_items
    elif remap_score_only or is_score_only:
        _assign_questions_to_catalog(question_items, catalog)

    if any(not item.get("co_code") for item in question_items):
        _fill_context_outcomes_from_catalog(context, catalog)
        _refresh_context_narrative(context)
        return context

    _attach_question_statistics(question_items, students)
    outcome_inputs = _merge_existing_outcome_inputs(context.get("course_outcomes") or [], catalog)
    course_outcomes = _course_outcomes_from_question_items(
        question_items,
        students,
        default_threshold=_catalog_default_threshold(catalog),
        outcome_inputs=outcome_inputs,
    )
    context.update(
        _build_context_from_parts(
            meta=context.get("meta", {}),
            students=students,
            question_items=question_items,
            course_outcomes=course_outcomes,
            input_mode=context.get("input_mode"),
        )
    )
    return context


def _merge_existing_outcome_inputs(
    existing_outcomes: list[dict[str, Any]],
    catalog: list[dict[str, Any]],
) -> dict[str, dict[str, Any]]:
    rows = {item["co_code"]: dict(item) for item in catalog}
    for index, item in enumerate(existing_outcomes, start=1):
        co_code = _normalize_co_code(item.get("co_code") or item.get("label") or f"CO{index}")
        row = rows.setdefault(
            co_code,
            {
                "co_code": co_code,
                "co_name": f"课程目标{_co_number(co_code) or index}",
                "label": f"课程目标{_co_number(co_code) or index}",
                "indicator": "",
                "description": "",
                "threshold": 0.65,
            },
        )
        for field in ("co_name", "label", "indicator", "description"):
            value = item.get(field)
            if value:
                row[field] = value
        threshold = _safe_float(item.get("threshold"))
        if threshold:
            row["threshold"] = threshold
    return rows


def _fill_context_outcomes_from_catalog(context: dict[str, Any], catalog: list[dict[str, Any]]) -> None:
    catalog_map = {item["co_code"]: item for item in catalog}
    for index, item in enumerate(context.get("course_outcomes") or [], start=1):
        co_code = _normalize_co_code(item.get("co_code") or item.get("label") or f"CO{index}")
        catalog_item = catalog_map.get(co_code)
        if not catalog_item:
            continue
        item["co_code"] = co_code
        item["co_name"] = item.get("co_name") or catalog_item.get("co_name") or catalog_item.get("label")
        item["label"] = item.get("label") or catalog_item.get("label") or item["co_name"]
        item["indicator"] = item.get("indicator") or catalog_item.get("indicator") or ""
        item["description"] = item.get("description") or catalog_item.get("description") or ""
        item["threshold"] = _safe_float(item.get("threshold")) or _safe_float(catalog_item.get("threshold")) or 0.65


def _refresh_context_narrative(context: dict[str, Any]) -> None:
    narrative_sections = _build_teacher_narrative_sections(
        meta=context.get("meta", {}),
        score_stats=context.get("score_stats", {}),
        score_segments=context.get("score_segments", []),
        question_groups=context.get("question_groups", []),
        course_outcomes=context.get("course_outcomes", []),
        student_outcomes=context.get("student_outcomes", []),
        overall_achievement=float(context.get("overall_achievement") or 0),
    )
    context["narrative_sections"] = narrative_sections
    context["narrative"] = _combine_narrative_sections(narrative_sections)


def _normalize_catalog_outcomes(outcomes: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for index, item in enumerate(outcomes, start=1):
        co_code = _normalize_co_code(item.get("co_code") or item.get("label") or f"CO{index}")
        rows.append(
            {
                "co_code": co_code,
                "co_name": str(item.get("co_name") or item.get("label") or f"课程目标{_co_number(co_code) or index}"),
                "label": str(item.get("label") or item.get("co_name") or f"课程目标{_co_number(co_code) or index}"),
                "indicator": str(item.get("indicator") or ""),
                "description": str(item.get("description") or ""),
                "threshold": float(item.get("threshold") or 0.65),
            }
        )
    return sorted(rows, key=lambda item: _co_sort_key(item["co_code"]))


def _assign_questions_to_catalog(question_items: list[dict[str, Any]], catalog: list[dict[str, Any]]) -> None:
    if not question_items or not catalog:
        return
    total_questions = len(question_items)
    total_outcomes = len(catalog)
    for index, item in enumerate(question_items):
        catalog_index = min(int(index * total_outcomes / total_questions), total_outcomes - 1)
        item["co_code"] = catalog[catalog_index]["co_code"]


def _synthesize_questions_from_final_scores(
    students: list[dict[str, Any]],
    catalog: list[dict[str, Any]],
    existing_questions: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    total_score = _simple_exam_total(existing_questions) or 100.0
    full_scores = _split_total_score(total_score, len(catalog))
    questions: list[dict[str, Any]] = []
    for index, (outcome, full_score) in enumerate(zip(catalog, full_scores), start=1):
        label = f"期末总分折算{outcome['co_code']}"
        for student in students:
            final_score = float(student.get("final_score") or 0)
            student.setdefault("question_scores", {})[label] = round(final_score * full_score / total_score, 2)
        questions.append(
            {
                "label": label,
                "qno": str(index),
                "qgroup_name": "期末总分折算",
                "full_score": round(full_score, 2),
                "co_code": outcome["co_code"],
                "avg_score": 0.0,
                "difficulty": 0.0,
                "pass_rate": 0.0,
            }
        )
    return questions


def _split_total_score(total_score: float, parts: int) -> list[float]:
    if parts <= 0:
        return []
    base = round(total_score / parts, 2)
    scores = [base for _ in range(parts)]
    scores[-1] = round(total_score - sum(scores[:-1]), 2)
    return scores


def _catalog_default_threshold(catalog: list[dict[str, Any]]) -> float:
    if not catalog:
        return 0.65
    return round(sum(float(item.get("threshold") or 0.65) for item in catalog) / len(catalog), 4)


def _parse_simple_input_workbook(workbook_value, exam_date_override: str | None = None) -> dict[str, Any]:
    info_ws = _get_sheet_by_normalized_name(workbook_value, "基本信息")
    structure_ws = _get_sheet_by_normalized_name(workbook_value, "试卷结构")
    outcomes_ws = _get_sheet_by_normalized_name(workbook_value, "课程目标")
    scores_ws = _get_sheet_by_normalized_name(workbook_value, "学生成绩")

    meta = _parse_simple_meta(info_ws, exam_date_override=exam_date_override)
    question_items = _parse_simple_question_items(structure_ws)
    students = _parse_simple_students(scores_ws, question_items)
    if not students:
        raise HTTPException(status_code=400, detail="学生成绩表至少需要填写一名学生")
    if not question_items:
        raise HTTPException(status_code=400, detail="试卷结构表至少需要填写一道题目")

    _attach_question_statistics(question_items, students)
    course_outcomes = _parse_simple_course_outcomes(
        outcomes_ws,
        question_items,
        students,
        default_threshold=float(meta.get("outcome_threshold") or 0.65),
    )
    meta["student_count_actual"] = len(students)
    if not meta.get("student_count_expected"):
        meta["student_count_expected"] = len(students)

    score_stats = _score_stats(students)
    score_segments = _score_segments(students)
    question_groups = _question_group_summary(students, question_items)
    student_outcomes = _simple_student_outcomes(students, question_items, course_outcomes)
    overall_achievement = _overall_achievement(student_outcomes)
    warnings = _warning_rows(students, overall_achievement)
    narrative_sections = _build_teacher_narrative_sections(
        meta=meta,
        score_stats=score_stats,
        score_segments=score_segments,
        question_groups=question_groups,
        course_outcomes=course_outcomes,
        student_outcomes=student_outcomes,
        overall_achievement=overall_achievement,
    )

    return {
        "meta": meta,
        "students": students,
        "score_stats": score_stats,
        "score_segments": score_segments,
        "difficulty_label": _difficulty_label(score_stats["average_score"], _simple_exam_total(question_items)),
        "question_items": question_items,
        "question_groups": question_groups,
        "course_outcomes": course_outcomes,
        "student_outcomes": student_outcomes,
        "overall_achievement": overall_achievement,
        "warnings": warnings,
        "narrative_sections": narrative_sections,
        "narrative": _combine_narrative_sections(narrative_sections),
    }


def _parse_simple_meta(ws, exam_date_override: str | None = None) -> dict[str, Any]:
    values: dict[str, str] = {}
    for row in range(1, ws.max_row + 1):
        key = _normalize_header(ws.cell(row, 1).value)
        value = _text(ws.cell(row, 2).value)
        if key and key not in {"字段", "项目"}:
            values[key] = value

    academic_year = _first_value(values, "学年学期", "学期", "学年")
    department = _first_value(values, "院系", "承担单位", "开课单位")
    course_code = _first_value(values, "课程代码", "课程编码", "课程编号")
    course_name = _first_value(values, "课程名称", "课程名")
    teacher_name = _first_value(values, "任课教师", "教师", "代课教师")
    class_name = _first_value(values, "班级", "教学班级", "专业班级")
    exam_date = (exam_date_override or "").strip() or _first_value(values, "考试日期", "考试时间")
    expected = _safe_float(_first_value(values, "学生人数", "应考人数"))
    outcome_threshold = _safe_float(_first_value(values, "课程目标达成阈值", "达成阈值", "默认阈值")) or 0.65

    return {
        "academic_year": academic_year,
        "department": department,
        "course_code": course_code,
        "course_name": course_name,
        "teacher_name": teacher_name,
        "class_name": class_name,
        "student_count_expected": int(expected) if expected else 0,
        "student_count_actual": 0,
        "exam_date": exam_date,
        "outcome_threshold": outcome_threshold,
    }


def _first_value(values: dict[str, str], *keys: str) -> str:
    for key in keys:
        value = values.get(_normalize_header(key), "")
        if value:
            return value
    return ""


def _parse_simple_question_items(ws) -> list[dict[str, Any]]:
    header = _header_map(ws, 1)
    required = ["题型", "满分", "课程目标编号"]
    if any(name not in header for name in required):
        raise HTTPException(status_code=400, detail="试卷结构表需要包含：题型、满分、课程目标编号")

    items: list[dict[str, Any]] = []
    used_labels: set[str] = set()
    for row in range(2, ws.max_row + 1):
        qtype = _text(ws.cell(row, header["题型"]).value)
        full_score = _safe_float(ws.cell(row, header["满分"]).value)
        co_code = _normalize_co_code(ws.cell(row, header["课程目标编号"]).value)
        if not qtype and full_score == 0 and not co_code:
            continue
        if not qtype or full_score <= 0 or not co_code:
            raise HTTPException(status_code=400, detail=f"试卷结构第 {row} 行缺少题型、满分或课程目标编号")
        qno = _text(ws.cell(row, header.get("题号", 0)).value) if header.get("题号") else str(len(items) + 1)
        sub_qno = _text(ws.cell(row, header.get("小题号", 0)).value) if header.get("小题号") else ""
        label = _simple_question_label(qtype, sub_qno, qno, used_labels)
        used_labels.add(label)
        items.append(
            {
                "label": label,
                "qno": qno or str(len(items) + 1),
                "qgroup_name": qtype,
                "full_score": round(full_score, 2),
                "co_code": co_code,
                "avg_score": 0.0,
                "difficulty": 0.0,
                "pass_rate": 0.0,
            }
        )
    return items


def _parse_simple_students(ws, question_items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    header = _header_map(ws, 1)
    if "学号" not in header:
        raise HTTPException(status_code=400, detail="学生成绩表需要包含：学号")

    question_columns = {}
    normalized_headers = {_normalize_header(ws.cell(1, col).value): col for col in range(1, ws.max_column + 1)}
    for item in question_items:
        column = normalized_headers.get(_normalize_header(item["label"]))
        if not column:
            raise HTTPException(status_code=400, detail=f"学生成绩表缺少题目得分列：{item['label']}")
        question_columns[item["label"]] = column

    final_score_col = None
    for name in ("期末总分", "总分", "成绩"):
        if name in header:
            final_score_col = header[name]
            break

    students: list[dict[str, Any]] = []
    for row in range(2, ws.max_row + 1):
        student_no = _text(ws.cell(row, header["学号"]).value)
        if not student_no:
            continue
        if not re.fullmatch(r"\d{6,20}(?:\.0)?", student_no):
            continue
        student_no = student_no[:-2] if student_no.endswith(".0") else student_no
        question_scores = {
            label: _safe_float(ws.cell(row, column).value)
            for label, column in question_columns.items()
        }
        computed_total = round(sum(question_scores.values()), 2)
        final_score = _safe_float(ws.cell(row, final_score_col).value) if final_score_col else computed_total
        students.append(
            {
                "student_no": student_no,
                "name": _text(ws.cell(row, header.get("姓名", 0)).value) if header.get("姓名") else "",
                "question_scores": question_scores,
                "final_score": round(final_score or computed_total, 2),
            }
        )
    return students


def _parse_simple_course_outcomes(
    ws,
    question_items: list[dict[str, Any]],
    students: list[dict[str, Any]],
    *,
    default_threshold: float = 0.65,
) -> list[dict[str, Any]]:
    header = _header_map(ws, 1)
    if "课程目标编号" not in header:
        raise HTTPException(status_code=400, detail="课程目标表需要包含：课程目标编号")

    outcome_inputs: dict[str, dict[str, Any]] = {}
    for row in range(2, ws.max_row + 1):
        co_code = _normalize_co_code(ws.cell(row, header["课程目标编号"]).value)
        if not co_code:
            continue
        threshold = _safe_float(ws.cell(row, header.get("达成阈值", 0)).value) if header.get("达成阈值") else 0
        outcome_inputs[co_code] = {
            "threshold": threshold or default_threshold,
            "indicator": _text(ws.cell(row, header.get("指标点", 0)).value) if header.get("指标点") else "",
            "description": _text(ws.cell(row, header.get("课程目标说明", 0)).value) if header.get("课程目标说明") else "",
        }

    co_codes = sorted({item["co_code"] for item in question_items}, key=_co_sort_key)
    total_full = _simple_exam_total(question_items) or 1.0
    rows: list[dict[str, Any]] = []
    for index, co_code in enumerate(co_codes, start=1):
        linked_items = [item for item in question_items if item["co_code"] == co_code]
        full_score = round(sum(float(item["full_score"]) for item in linked_items), 2)
        scores = [
            sum(float(student["question_scores"].get(item["label"], 0)) for item in linked_items)
            for student in students
        ]
        avg_score = round(sum(scores) / len(scores), 2) if scores else 0.0
        achievement = round(avg_score / full_score, 4) if full_score else 0.0
        input_row = outcome_inputs.get(co_code, {})
        label = f"课程目标{_co_number(co_code) or index}"
        supporting_groups: list[str] = []
        for item in linked_items:
            group_name = str(item["qgroup_name"])
            if group_name not in supporting_groups:
                supporting_groups.append(group_name)
        weight = round(full_score / total_full, 4)
        rows.append(
            {
                "label": label,
                "co_code": co_code,
                "co_name": label,
                "indicator": input_row.get("indicator", ""),
                "description": input_row.get("description", ""),
                "full_score": full_score,
                "avg_score": avg_score,
                "achievement": achievement,
                "threshold": round(float(input_row.get("threshold") or default_threshold), 4),
                "supporting_groups": supporting_groups,
                "weight": weight,
                "result": round(weight * achievement, 4),
            }
        )
    return rows


def _simple_student_outcomes(
    students: list[dict[str, Any]],
    question_items: list[dict[str, Any]],
    course_outcomes: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    total_full = _simple_exam_total(question_items) or 1.0
    for student in students:
        overall = round(float(student["final_score"]) / total_full, 4) if total_full else 0.0
        for outcome in course_outcomes:
            linked_items = [item for item in question_items if item["co_code"] == outcome["co_code"]]
            co_score = round(sum(float(student["question_scores"].get(item["label"], 0)) for item in linked_items), 2)
            full_score = float(outcome.get("full_score") or 0)
            achievement = round(co_score / full_score, 4) if full_score else 0.0
            threshold = float(outcome.get("threshold") or 0.65)
            rows.append(
                {
                    "student_no": student["student_no"],
                    "label": outcome["label"],
                    "co_code": outcome["co_code"],
                    "co_score": co_score,
                    "achievement": achievement,
                    "full_score": round(full_score, 2),
                    "threshold": round(threshold, 4),
                    "is_attained": achievement >= threshold,
                    "overall_achievement": overall,
                }
            )
    return rows


def _attach_question_statistics(question_items: list[dict[str, Any]], students: list[dict[str, Any]]) -> None:
    for item in question_items:
        scores = [float(student["question_scores"].get(item["label"], 0)) for student in students]
        full_score = float(item.get("full_score") or 0)
        avg_score = sum(scores) / len(scores) if scores else 0.0
        item["avg_score"] = round(avg_score, 2)
        item["difficulty"] = round(avg_score / full_score, 4) if full_score else 0.0
        item["pass_rate"] = round(sum(score >= full_score * 0.6 for score in scores) / len(scores), 4) if scores else 0.0


def _header_map(ws, row: int) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for column in range(1, ws.max_column + 1):
        name = _normalize_header(ws.cell(row, column).value)
        if name:
            mapping[name] = column
    return mapping


def _normalize_header(value: Any) -> str:
    return re.sub(r"\s+", "", _text(value))


def _normalize_co_code(value: Any) -> str:
    text = _normalize_header(value).upper()
    if not text:
        return ""
    match = re.search(r"(?:CO|课程目标)?(\d+)", text, flags=re.IGNORECASE)
    return f"CO{int(match.group(1))}" if match else text


def _co_sort_key(co_code: str) -> tuple[int, str]:
    number = _co_number(co_code)
    return (number or 9999, co_code)


def _co_number(co_code: Any) -> int | None:
    match = re.search(r"(\d+)", str(co_code or ""))
    return int(match.group(1)) if match else None


def _simple_question_label(qtype: str, sub_qno: str, qno: str, used_labels: set[str]) -> str:
    base = f"{qtype}{sub_qno}" if sub_qno else qtype
    if base not in used_labels:
        return base
    fallback = f"{qtype}{qno}" if qno else f"{qtype}{len(used_labels) + 1}"
    return fallback if fallback not in used_labels else f"{fallback}_{len(used_labels) + 1}"


def _simple_exam_total(question_items: list[dict[str, Any]]) -> float:
    return round(sum(float(item.get("full_score") or 0) for item in question_items), 2)


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _safe_float(value: Any) -> float:
    try:
        if value is None or value == "":
            return 0.0
        return round(float(value), 4)
    except Exception:
        return 0.0


def _parse_meta(raw_ws, main_ws, exam_date_override: str | None = None) -> dict[str, Any]:
    academic_year = (_text(main_ws["A2"].value) or _text(raw_ws["A1"].value).replace("末成绩单", "")).replace("-", "—")
    department_course = _text(main_ws["A3"].value)
    teacher_class = _text(main_ws["A4"].value)
    raw_class_line = _text(raw_ws["A3"].value)

    department = _extract(r"承担单位[:：]\s*(.+?)\s+课程名称", department_course)
    course_name = _extract(r"课程名称[:：]\s*(.+)$", department_course) or _extract(r"课程名称[:：]\s*(.+)$", _text(raw_ws["A2"].value))
    teacher_name = _extract(r"任课教师[:：]\s*(.+?)\s+班级", teacher_class) or _extract(r"代课教师[:：]\s*(.+)$", _text(raw_ws["F2"].value))
    class_name = _extract(r"班级[:：]\s*(.+)$", teacher_class) or _extract(r"专业班级[:：]\s*(.+?)\s+学生人数", raw_class_line)
    student_count_expected = int(_extract_number(r"学生人数[:：]\s*(\d+)", raw_class_line) or 0)
    exam_date = (exam_date_override or "").strip() or _extract(r"考试时间[:：]\s*(.+)$", teacher_class) or ""

    return {
        "academic_year": academic_year,
        "department": department,
        "course_name": course_name,
        "teacher_name": teacher_name,
        "class_name": class_name,
        "student_count_expected": student_count_expected,
        "student_count_actual": 0,
        "exam_date": exam_date,
    }


def _extract(pattern: str, text: str) -> str:
    match = re.search(pattern, text)
    return match.group(1).strip() if match else ""


def _extract_number(pattern: str, text: str) -> int | None:
    match = re.search(pattern, text)
    return int(match.group(1)) if match else None


def _resolve_merged_value(ws, row: int, column: int) -> Any:
    value = ws.cell(row, column).value
    if value is not None:
        return value
    for cell_range in ws.merged_cells.ranges:
        if cell_range.min_row <= row <= cell_range.max_row and cell_range.min_col <= column <= cell_range.max_col:
            return ws.cell(cell_range.min_row, cell_range.min_col).value
    return None


def _parse_students_and_questions(raw_ws, main_ws) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    data_start_row = 6
    data_end_row = data_start_row
    while _text(raw_ws.cell(data_end_row, 1).value):
        data_end_row += 1
    data_end_row -= 1

    question_items: list[dict[str, Any]] = []
    question_columns: list[dict[str, Any]] = []
    for column in range(3, 12):
        group_name = _text(_resolve_merged_value(raw_ws, 4, column) or _resolve_merged_value(main_ws, 5, column))
        sub_label = _text(raw_ws.cell(5, column).value)
        full_score = _safe_float(main_ws.cell(8, column).value)
        if not group_name or group_name in {"成绩", "总分"}:
            continue
        if sub_label in {"得分", "总分"}:
            continue
        item_label = f"{group_name}{sub_label}" if sub_label else group_name
        question_columns.append(
            {
                "column": column,
                "group_name": group_name,
                "label": item_label,
                "full_score": full_score,
            }
        )

    students: list[dict[str, Any]] = []
    for row in range(data_start_row, data_end_row + 1):
        student_no = _text(raw_ws.cell(row, 1).value)
        if not student_no or not re.fullmatch(r"\d{6,20}", student_no):
            continue
        question_scores: dict[str, float] = {}
        for item in question_columns:
            question_scores[item["label"]] = _safe_float(raw_ws.cell(row, item["column"]).value)
        students.append(
            {
                "student_no": student_no,
                "name": _text(raw_ws.cell(row, 2).value),
                "question_scores": question_scores,
                "final_score": _safe_float(raw_ws.cell(row, 12).value),
            }
        )

    if students:
        averages = defaultdict(list)
        for student in students:
            for label, score in student["question_scores"].items():
                averages[label].append(score)
        for item in question_columns:
            scores = averages[item["label"]]
            avg_score = sum(scores) / len(scores) if scores else 0.0
            pass_rate = sum(score >= item["full_score"] * 0.6 for score in scores) / len(scores) if scores else 0.0
            question_items.append(
                {
                    "label": item["label"],
                    "qgroup_name": item["group_name"],
                    "full_score": round(item["full_score"], 2),
                    "avg_score": round(avg_score, 2),
                    "difficulty": round(avg_score / item["full_score"], 4) if item["full_score"] else 0.0,
                    "pass_rate": round(pass_rate, 4),
                }
            )

    return students, question_items


def _score_stats(students: list[dict[str, Any]]) -> dict[str, Any]:
    if not students:
        return {"total_students": 0, "average_score": 0.0, "max_score": 0.0, "min_score": 0.0, "pass_rate": 0.0}
    scores = [item["final_score"] for item in students]
    total = len(scores)
    return {
        "total_students": total,
        "average_score": round(sum(scores) / total, 2),
        "max_score": round(max(scores), 2),
        "min_score": round(min(scores), 2),
        "pass_rate": round(sum(score >= 60 for score in scores) / total, 4),
    }


def _score_segments(students: list[dict[str, Any]]) -> list[dict[str, Any]]:
    scores = [item["final_score"] for item in students]
    total = len(scores) or 1
    buckets = [
        ("90—100分", lambda score: 90 <= score <= 100),
        ("80—89分", lambda score: 80 <= score < 90),
        ("70—79分", lambda score: 70 <= score < 80),
        ("60—69分", lambda score: 60 <= score < 70),
        ("50—59分", lambda score: 50 <= score < 60),
        ("50分以下", lambda score: score < 50),
    ]
    rows = []
    for label, rule in buckets:
        count = sum(1 for score in scores if rule(score))
        rows.append({"label": label, "count": count, "rate": round(count / total, 4)})
    return rows


def _difficulty_label(average_score: float, total_score: float) -> str:
    if total_score <= 0:
        return "未知"
    ratio = average_score / total_score
    if ratio >= 0.8:
        return "容易"
    if ratio >= 0.6:
        return "中等"
    return "较难"


def _question_group_summary(students: list[dict[str, Any]], question_items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not students or not question_items:
        return []
    full_map = defaultdict(float)
    avg_map = defaultdict(float)
    for item in question_items:
        full_map[item["qgroup_name"]] += item["full_score"]
        avg_map[item["qgroup_name"]] += item["avg_score"]
    rows = []
    for group_name, full_score in full_map.items():
        avg_score = avg_map[group_name]
        rows.append(
            {
                "qgroup_name": group_name,
                "full_score": round(full_score, 2),
                "avg_score": round(avg_score, 2),
                "achievement": round(avg_score / full_score, 4) if full_score else 0.0,
            }
        )
    return rows


def _threshold_map(chart_ws) -> dict[str, float]:
    if chart_ws is None:
        return {}
    mapping: dict[str, float] = {}
    for row in range(2, chart_ws.max_row + 1):
        label = _text(chart_ws.cell(row, 1).value)
        threshold = _safe_float(chart_ws.cell(row, 2).value)
        if label:
            mapping[label] = threshold or 0.65
    return mapping


def _supporting_groups(main_ws_formula, main_ws_value) -> dict[str, list[str]]:
    mapping: dict[str, list[str]] = {}
    formula_cells = {"课程目标1": "M8", "课程目标2": "O8", "课程目标3": "Q8", "课程目标4": "S8"}
    for label, cell_ref in formula_cells.items():
        formula = _text(main_ws_formula[cell_ref].value)
        columns = re.findall(r"([A-Z]+)8", formula)
        groups: list[str] = []
        for column_name in columns:
            column_index = column_index_from_string(column_name)
            group_name = _text(_resolve_merged_value(main_ws_value, 5, column_index))
            if group_name in {"得分", "总分"}:
                group_name = _nearest_left_group_name(main_ws_value, column_index)
            if group_name and group_name not in groups:
                groups.append(group_name)
        mapping[label] = groups
    return mapping


def _nearest_left_group_name(ws, column_index: int) -> str:
    for current_column in range(column_index - 1, 2, -1):
        group_name = _text(_resolve_merged_value(ws, 5, current_column))
        if group_name and group_name not in {"得分", "总分"}:
            return group_name
    return ""


def _course_outcomes(summary_ws, threshold_map: dict[str, float], support_map: dict[str, list[str]]) -> list[dict[str, Any]]:
    markers = []
    for row in range(2, summary_ws.max_row + 1):
        label = _text(summary_ws.cell(row, 1).value)
        if label.startswith("课程目标"):
            markers.append((label, row))
    markers.append(("", summary_ws.max_row + 1))

    rows: list[dict[str, Any]] = []
    for index in range(len(markers) - 1):
        label, start_row = markers[index]
        end_row = markers[index + 1][1] - 1
        summary_row = None
        for row in range(start_row, end_row + 1):
            full_score = summary_ws.cell(row, 3).value
            avg_score = summary_ws.cell(row, 4).value
            achievement = summary_ws.cell(row, 5).value
            if full_score is not None and avg_score is not None and achievement is not None:
                summary_row = row
        if summary_row is None:
            continue
        full_score = _safe_float(summary_ws.cell(summary_row, 3).value)
        avg_score = _safe_float(summary_ws.cell(summary_row, 4).value)
        achievement = _safe_float(summary_ws.cell(summary_row, 5).value)
        co_index = _extract_number(r"(\d+)", label) or len(rows) + 1
        rows.append(
            {
                "label": label,
                "co_code": f"CO{co_index}",
                "full_score": round(full_score, 2),
                "avg_score": round(avg_score, 2),
                "achievement": round(achievement, 4),
                "threshold": round(threshold_map.get(label, 0.65), 4),
                "supporting_groups": support_map.get(label, []),
            }
        )

    total_full_score = sum(item["full_score"] for item in rows) or 1.0
    for item in rows:
        item["weight"] = round(item["full_score"] / total_full_score, 4)
        item["result"] = round(item["weight"] * item["achievement"], 4)
    return rows


def _student_outcomes(main_ws, threshold_map: dict[str, float]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    data_row = 9
    while _text(main_ws.cell(data_row, 1).value):
        student_no = _text(main_ws.cell(data_row, 1).value)
        for index, (score_col, achievement_col) in enumerate(((13, 14), (15, 16), (17, 18), (19, 20)), start=1):
            label = f"课程目标{index}"
            threshold = threshold_map.get(label, 0.65)
            co_score = _safe_float(main_ws.cell(data_row, score_col).value)
            achievement = _safe_float(main_ws.cell(data_row, achievement_col).value)
            full_score = _safe_float(main_ws.cell(8, score_col).value)
            rows.append(
                {
                    "student_no": student_no,
                    "label": label,
                    "co_code": f"CO{index}",
                    "co_score": round(co_score, 2),
                    "achievement": round(achievement, 4),
                    "full_score": round(full_score, 2),
                    "threshold": round(threshold, 4),
                    "is_attained": achievement >= threshold,
                    "overall_achievement": _safe_float(main_ws.cell(data_row, 21).value),
                }
            )
        data_row += 1
    return rows


def _overall_achievement(student_outcomes: list[dict[str, Any]]) -> float:
    values = [item["overall_achievement"] for item in student_outcomes if item["overall_achievement"] > 0]
    if not values:
        return 0.0
    return round(sum(values) / len(values), 4)


def _warning_rows(students: list[dict[str, Any]], overall_achievement: float) -> list[dict[str, Any]]:
    warnings = []
    for item in students:
        reasons: list[str] = []
        if item["final_score"] < 60:
            reasons.append("期末卷面成绩低于60分")
        if item["final_score"] < 50:
            reasons.append("低分风险明显")
        if reasons:
            warnings.append({"student_no": item["student_no"], "reasons": reasons, "final_score": item["final_score"]})
    if not warnings and overall_achievement < 0.65:
        warnings.append({"student_no": "-", "reasons": ["班级整体课程目标达成度低于阈值"], "final_score": 0})
    return warnings


def _build_teacher_narrative_sections(
    *,
    meta: dict[str, Any],
    score_stats: dict[str, Any],
    score_segments: list[dict[str, Any]],
    question_groups: list[dict[str, Any]],
    course_outcomes: list[dict[str, Any]],
    student_outcomes: list[dict[str, Any]],
    overall_achievement: float,
) -> dict[str, str]:
    intro = [
        f"本次《{meta['course_name']}》考核结果来自成绩工作簿，分析识别出 {len(question_groups)} 类主要题型，"
        f"班级共 {score_stats['total_students']} 人，平均分 {score_stats['average_score']} 分，试题整体难度判断为{_difficulty_label(score_stats['average_score'], 100.0)}。",
    ]
    for item in course_outcomes:
        reached = sum(
            1
            for row in student_outcomes
            if row["label"] == item["label"] and row["achievement"] >= row["threshold"]
        )
        support_text = "、".join(item["supporting_groups"]) if item["supporting_groups"] else "相关题型"
        intro.append(
            f"{item['label']}主要通过{support_text}进行考查，分值为{int(item['full_score'])}分，"
            f"平均分为{item['avg_score']}分，达成度为{_format_ratio(item['achievement'])}，"
            f"阈值为{_format_ratio(item['threshold'])}，达到阈值的学生有{reached}人。"
        )

    middle = [
        f"根据工作簿中的学生成绩与课程目标汇总结果，班级整体课程目标达成度约为{_format_ratio(overall_achievement)}。",
    ]
    weakest = [item for item in course_outcomes if item["achievement"] < item["threshold"]]
    if weakest:
        middle.append(
            "当前相对薄弱的目标为"
            + "、".join(f"{item['label']}（{_format_ratio(item['achievement'])}）" for item in weakest)
            + "，后续应重点围绕这些目标组织讲评与训练。"
        )
    else:
        middle.append("各课程目标均达到预设阈值，说明学生整体作答情况能够支撑课程目标达成。")

    dominant_segment = max(score_segments, key=lambda item: item["count"]) if score_segments else {"label": "无数据", "count": 0}
    ending = [
        f"成绩主要集中在{dominant_segment['label']}，共{dominant_segment['count']}人。建议继续保留行之有效的知识点讲解方式，"
        "并针对低分题型与未达标课程目标增加分层讲评、错题复盘和案例训练。",
    ]
    if weakest:
        for item in weakest:
            ending.append(
                f"围绕{item['label']}，建议强化{('、'.join(item['supporting_groups']) or '对应题型')}的专题训练，"
                "通过阶段性测验和可视化反馈帮助学生及时纠偏。"
            )
    else:
        ending.append("在整体达标的基础上，可通过综合性实验和项目训练继续提升高阶应用能力。")

    return {
        "score_summary": (
            f"本次《{meta['course_name']}》共纳入 {score_stats['total_students']} 名学生，平均分 {score_stats['average_score']} 分，"
            f"最高分 {score_stats['max_score']} 分，最低分 {score_stats['min_score']} 分，及格率 {_format_percent(score_stats['pass_rate'])}。"
            f"试卷整体难度判断为{_difficulty_label(score_stats['average_score'], 100.0)}。"
        ),
        "support_analysis": "".join(intro),
        "attainment_analysis": "".join(middle),
        "improvement_actions": "".join(ending),
    }


def _build_teacher_narrative(
    *,
    meta: dict[str, Any],
    score_stats: dict[str, Any],
    score_segments: list[dict[str, Any]],
    question_groups: list[dict[str, Any]],
    course_outcomes: list[dict[str, Any]],
    student_outcomes: list[dict[str, Any]],
    overall_achievement: float,
) -> str:
    sections = _build_teacher_narrative_sections(
        meta=meta,
        score_stats=score_stats,
        score_segments=score_segments,
        question_groups=question_groups,
        course_outcomes=course_outcomes,
        student_outcomes=student_outcomes,
        overall_achievement=overall_achievement,
    )
    return _combine_narrative_sections(sections)


def _format_ratio(value: float) -> str:
    return f"{round(value, 2):.2f}".rstrip("0").rstrip(".")


def _build_report_document(contexts: list[dict[str, Any]]) -> Document:
    if not TEMPLATE_DOCX_PATH.exists():
        raise HTTPException(status_code=500, detail="报告样式文件缺失")

    _configure_chart_font()
    document = Document(TEMPLATE_DOCX_PATH)
    body = document._element.body
    children = list(body)
    for child in children[3:-1]:
        body.remove(child)

    _fill_report_header(document.paragraphs[:2], contexts[0]["meta"]["academic_year"])
    _fill_report_table(document.tables[0], contexts[0])

    for context in contexts[1:]:
        document.add_page_break()
        paragraph_1 = document.add_paragraph()
        paragraph_1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_1 = paragraph_1.add_run(f"吕梁学院{context['meta']['academic_year']}")
        run_1.font.size = Pt(16)
        run_1.bold = False

        paragraph_2 = document.add_paragraph()
        paragraph_2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_2 = paragraph_2.add_run("试 卷 分 析 表")
        run_2.font.size = Pt(18)
        run_2.bold = False

        template_table = copy.deepcopy(document.tables[0]._element)
        body.insert_element_before(template_table, "w:sectPr")
        _fill_report_table(document.tables[-1], context)

    return document


def _fill_report_header(paragraphs, academic_year: str) -> None:
    if len(paragraphs) >= 2:
        paragraphs[0].text = f"吕梁学院{academic_year}"
        paragraphs[1].text = "试 卷 分 析 表"
        for paragraph in paragraphs[:2]:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _fill_report_table(table, context: dict[str, Any]) -> None:
    meta = context["meta"]
    stats = context["score_stats"]
    segments = {item["label"]: item for item in context["score_segments"]}

    _set_cell(table.cell(0, 2), meta["department"] or DEFAULT_DEPARTMENT)
    _set_cell(table.cell(0, 5), meta["class_name"] or "待补充")
    _set_cell(table.cell(0, 7), meta["course_name"] or "待补充")
    _set_cell(table.cell(1, 2), str(meta["student_count_expected"] or stats["total_students"]))
    _set_cell(table.cell(1, 5), str(meta["student_count_actual"] or stats["total_students"]))
    _set_cell(table.cell(1, 7), meta["exam_date"] or "待补充")

    _set_cell(table.cell(2, 1), _segment_text(segments.get("90—100分")))
    _set_cell(table.cell(2, 5), _segment_text(segments.get("60—69分")))
    _set_cell(table.cell(3, 1), _segment_text(segments.get("80—89分")))
    _set_cell(table.cell(3, 5), _segment_text(segments.get("50—59分")))
    _set_cell(table.cell(4, 1), _segment_text(segments.get("70—79分")))
    _set_cell(table.cell(4, 5), _segment_text(segments.get("50分以下")))

    _set_cell(table.cell(5, 3), context["difficulty_label"])
    _set_cell(table.cell(5, 6), _format_number(stats["average_score"]))
    _set_cell(table.cell(6, 3), _format_number(stats["max_score"]))
    _set_cell(table.cell(6, 6), _format_number(stats["min_score"]))

    _fill_analysis_cell(table.cell(8, 1), context)


def _segment_text(item: dict[str, Any] | None) -> str:
    if not item:
        return ""
    return f"{item['label']}{item['count']}人   占{_format_percent(item['rate'])}"


def _format_percent(value: float) -> str:
    return f"{round(value * 100, 2):.2f}%".replace(".00%", "%")


def _format_number(value: float) -> str:
    rounded = round(value, 2)
    if abs(rounded - int(rounded)) < 1e-9:
        return str(int(rounded))
    return f"{rounded:.2f}".rstrip("0").rstrip(".")


def _set_cell(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    run = paragraph.add_run(str(text))
    run.font.size = Pt(10.5)


def _set_multiline_cell(cell, text: str, font_size: float = 10.5) -> None:
    first_paragraph = _clear_cell(cell)
    lines = str(text).splitlines() or [""]
    for index, line in enumerate(lines):
        paragraph = first_paragraph if index == 0 else cell.add_paragraph()
        run = paragraph.add_run(line)
        run.font.size = Pt(font_size)


def _clear_cell(cell):
    tc = cell._tc
    tc_pr = tc.tcPr
    for child in list(tc):
        if child is not tc_pr:
            tc.remove(child)
    return cell.add_paragraph()


def _add_cell_paragraph(cell, text: str = "", *, font_size: float = 10.5, bold: bool = False, align=None):
    paragraph = cell.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    run.font.size = Pt(font_size)
    run.bold = bold
    return paragraph


def _set_table_cell_text(cell, text: str, *, font_size: float = 8.5, bold: bool = False) -> None:
    paragraph = _clear_cell(cell)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(str(text))
    run.font.size = Pt(font_size)
    run.bold = bold


def _fill_analysis_cell(cell, context: dict[str, Any]) -> None:
    _clear_cell(cell)
    profile = _course_goal_profile(context)

    _add_cell_paragraph(cell, "1、试题对课程目标的支撑度分析", font_size=11, bold=True)
    for paragraph_text in _support_analysis_paragraphs(context, profile):
        _add_cell_paragraph(cell, paragraph_text, font_size=10.5)

    _add_cell_paragraph(cell, "2、学生作答情况对课程目标的达成度分析", font_size=11, bold=True)
    _add_cell_paragraph(cell, _attainment_intro(context), font_size=10.5)
    title = _add_cell_paragraph(cell, "课程目标达成度表", font_size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    title.paragraph_format.space_before = Pt(4)
    _add_outcome_table(cell, context)
    _add_cell_picture(cell, _render_outcome_chart(context), width=5.8)
    if context.get("student_outcomes"):
        _add_cell_picture(cell, _render_student_outcome_distribution_chart(context), width=5.8)
    for item in context.get("course_outcomes") or []:
        _add_cell_paragraph(cell, _outcome_detail_paragraph(context, item, profile), font_size=10.5)

    _add_cell_paragraph(cell, "3、教师对今后教学持续改进的具体意见", font_size=11, bold=True)
    narrative_improvements = _narrative_improvement_paragraphs(context)
    if narrative_improvements:
        for paragraph_text in narrative_improvements:
            _add_cell_paragraph(cell, paragraph_text, font_size=10.5)
    else:
        for index, item in enumerate(context.get("course_outcomes") or [], start=1):
            goal_profile = profile[min(index - 1, len(profile) - 1)] if profile else _generic_goal_profile(index)
            _add_cell_paragraph(cell, f"（{index}）{goal_profile['improvement_title']}", font_size=10.5, bold=True)
            _add_cell_paragraph(cell, _improvement_paragraph(context, item, goal_profile), font_size=10.5)
        _add_cell_paragraph(
            cell,
            "这些方法和策略有助于教师全面了解学生的学习情况，有效指导教学实践，提升教学质量和学生学习成效。",
            font_size=10.5,
        )

    _add_cell_paragraph(cell, "", font_size=10.5)
    _add_cell_paragraph(cell, "任课教师签名：", font_size=10.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _add_cell_paragraph(cell, "年    月    日", font_size=10.5, align=WD_ALIGN_PARAGRAPH.RIGHT)


def _narrative_improvement_paragraphs(context: dict[str, Any]) -> list[str]:
    sections = context.get("narrative_sections") or {}
    text = str(sections.get("improvement_actions") or "").strip()
    if not text:
        return []
    lines = [line.strip() for line in re.split(r"[\r\n]+", text) if line.strip()]
    return lines or [text]


def _add_outcome_table(cell, context: dict[str, Any]) -> None:
    outcomes = list(context.get("course_outcomes") or [])
    if not outcomes:
        _add_cell_paragraph(cell, "暂无课程目标达成度数据。", font_size=10.5)
        return

    table = cell.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    headers = ["课程目标", "分值", "", "总分", "平均分", "分项达成度", "权重", "达成结果"]
    table.cell(0, 1).merge(table.cell(0, 2))
    for index, title in enumerate(headers):
        if index == 2:
            continue
        paragraph = table.cell(0, index).paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(title)
        run.bold = True
        run.font.size = Pt(9)

    total_full_score = 0.0
    total_result = 0.0
    for item in outcomes:
        support_items = _support_items_for_outcome(context, item)
        support_items = support_items or [{"name": "未绑定题型", "score": float(item.get("full_score") or 0)}]
        first_row_index = len(table.rows)
        for support_index, support_item in enumerate(support_items):
            row_cells = table.add_row().cells
            values = [
                _outcome_label(item),
                support_item["name"],
                _format_number(float(support_item.get("score") or 0)),
                _format_number(float(item.get("full_score") or 0)),
                _format_number(float(item.get("avg_score") or 0)),
                _format_ratio(float(item.get("achievement") or 0)),
                _format_ratio(float(item.get("weight") or 0)),
                _format_ratio(float(item.get("result") or 0)),
            ]
            if support_index > 0:
                for merged_column in (0, 3, 4, 5, 6, 7):
                    values[merged_column] = ""
            for index, value in enumerate(values):
                paragraph = row_cells[index].paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(value)
                run.font.size = Pt(8.5)
        if len(support_items) > 1:
            last_row_index = len(table.rows) - 1
            merged_values = {
                0: _outcome_label(item),
                3: _format_number(float(item.get("full_score") or 0)),
                4: _format_number(float(item.get("avg_score") or 0)),
                5: _format_ratio(float(item.get("achievement") or 0)),
                6: _format_ratio(float(item.get("weight") or 0)),
                7: _format_ratio(float(item.get("result") or 0)),
            }
            for column_index, merged_value in merged_values.items():
                merged_cell = table.cell(first_row_index, column_index).merge(table.cell(last_row_index, column_index))
                _set_table_cell_text(merged_cell, merged_value)
        total_full_score += float(item.get("full_score") or 0)
        total_result += float(item.get("result") or 0)

    row_cells = table.add_row().cells
    row_cells[0].merge(row_cells[2])
    total_avg_score = float((context.get("score_stats") or {}).get("average_score") or 0)
    totals = ["合计", "", "", _format_number(total_full_score), _format_number(total_avg_score), "", "", _format_ratio(total_result)]
    for index, value in enumerate(totals):
        if index in {1, 2}:
            continue
        paragraph = row_cells[index].paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value)
        run.bold = index == 0
        run.font.size = Pt(8.5)


def _course_goal_profile(context: dict[str, Any]) -> list[dict[str, str]]:
    outcomes = context.get("course_outcomes") or []
    if outcomes and any(item.get("description") or item.get("indicator") for item in outcomes):
        return [
            {
                "indicator": str(item.get("indicator") or f"指标点{index}"),
                "description": str(
                    item.get("description")
                    or "能够围绕课程核心知识、方法和应用场景完成相应学习任务，支撑课程目标达成"
                ),
                "focus": str(item.get("description") or "课程核心知识与应用能力")[:40],
                "improvement_title": "强化课程目标对应能力训练",
                "improvement": "应结合该课程目标对应题型开展专题讲评、错题复盘和阶段性检测，帮助学生补齐薄弱环节。",
            }
            for index, item in enumerate(outcomes, start=1)
        ]
    course_name = str((context.get("meta") or {}).get("course_name") or "")
    for key, value in COURSE_GOAL_DESCRIPTIONS.items():
        if key in course_name:
            return value
    return [_generic_goal_profile(index) for index, _ in enumerate(outcomes, start=1)]


def _generic_goal_profile(index: int) -> dict[str, str]:
    return {
        "indicator": f"指标点{index}",
        "description": "能够围绕课程核心知识、方法和应用场景完成相应学习任务，支撑课程目标达成",
        "focus": "课程核心知识与应用能力",
        "improvement_title": "强化课程目标对应能力训练",
        "improvement": "应结合该课程目标对应题型开展专题讲评、错题复盘和阶段性检测，帮助学生补齐薄弱环节。",
    }


def _support_analysis_paragraphs(context: dict[str, Any], profile: list[dict[str, str]]) -> list[str]:
    meta = context.get("meta") or {}
    course_name = str(meta.get("course_name") or "本课程")
    group_names = _ordered_question_group_names(context)
    group_text = "、".join(group_names) if group_names else "相关题型"
    paragraphs = [
        f"本次《{course_name}》考核方式为笔试，分{group_text}等题型。《{course_name}》试题考察学生识记、理解、应用、分析、评价和创造的能力。"
    ]
    for index, item in enumerate(context.get("course_outcomes") or [], start=1):
        goal_profile = profile[min(index - 1, len(profile) - 1)] if profile else _generic_goal_profile(index)
        support_items = _support_items_for_outcome(context, item)
        support_text = _support_names_text(support_items)
        paragraphs.append(
            f"{_outcome_label(item)}：{goal_profile['description']}，实现对{goal_profile['indicator']}的支撑。"
            f"试题中{support_text}实现对{_outcome_label(item)}的支撑，分值为{_format_number(float(item.get('full_score') or 0))}分。"
        )
    return paragraphs


def _attainment_intro(context: dict[str, Any]) -> str:
    meta = context.get("meta") or {}
    outcomes = context.get("course_outcomes") or []
    threshold = _default_outcome_threshold(context)
    labels = "、".join(str(_outcome_number(item)) for item in outcomes) or "各"
    values = "、".join(_format_ratio(float(item.get("achievement") or 0)) for item in outcomes) or "暂无"
    return (
        f"依据《{meta.get('course_name') or '本课程'}》教学大纲，"
        f"课程目标达成度期望值为{_format_ratio(threshold)}，通过统计学生的得分情况，"
        f"课程目标{labels}的达成度分别为{values}，课程目标达成度为{_format_ratio(_overall_result(context))}。"
    )


def _outcome_detail_paragraph(context: dict[str, Any], item: dict[str, Any], profile: list[dict[str, str]]) -> str:
    index = max(_outcome_number(item), 1)
    goal_profile = profile[min(index - 1, len(profile) - 1)] if profile else _generic_goal_profile(index)
    support_items = _support_items_for_outcome(context, item)
    support_text = _support_names_text(support_items)
    threshold = float(item.get("threshold") or _default_outcome_threshold(context))
    attained_count = _outcome_attained_count(context, item)
    total_students = int((context.get("score_stats") or {}).get("total_students") or 0)
    support_counts = _support_score_count_text(context, support_items)
    judgement = _achievement_judgement(float(item.get("achievement") or 0), threshold)
    return (
        f"{_outcome_label(item)}主要通过{support_text}进行考查，总分值为{_format_number(float(item.get('full_score') or 0))}分，"
        f"学生平均得分为{_format_number(float(item.get('avg_score') or 0))}分，对应达成评价值为{_format_ratio(float(item.get('achievement') or 0))}，"
        f"期望值为{_format_ratio(threshold)}。整体来看，{support_counts}"
        f"{_outcome_label(item)}达成度在{_format_ratio(threshold)}以上的有{attained_count}人"
        f"{'，占' + _format_percent(attained_count / total_students) if total_students else ''}。"
        f"该目标主要检验学生{goal_profile['focus']}，作答结果显示{judgement}。"
        f"后续教学应围绕{support_text}继续开展针对性讲评和训练，以更好支撑{goal_profile['indicator']}的达成。"
    )


def _improvement_paragraph(context: dict[str, Any], item: dict[str, Any], goal_profile: dict[str, str]) -> str:
    support_text = _support_names_text(_support_items_for_outcome(context, item))
    return (
        f"{_outcome_label(item)}的整体达成度为{_format_ratio(float(item.get('achievement') or 0))}，"
        f"学生在{support_text}中反映出{goal_profile['focus']}方面仍需进一步巩固。"
        f"{goal_profile['improvement']}"
    )


def _ordered_question_group_names(context: dict[str, Any]) -> list[str]:
    names: list[str] = []
    for item in context.get("question_items") or []:
        name = _clean_group_name(item.get("qgroup_name") or item.get("qtype") or item.get("label"))
        if name and name not in names:
            names.append(name)
    for item in context.get("question_groups") or []:
        name = _clean_group_name(item.get("qgroup_name") or item.get("label"))
        if name and name not in names:
            names.append(name)
    return names


def _support_items_for_outcome(context: dict[str, Any], item: dict[str, Any]) -> list[dict[str, Any]]:
    co_code = str(item.get("co_code") or "")
    group_scores: dict[str, float] = {}
    for record in _unique_question_records(context):
        if str(record.get("co_code") or "") != co_code:
            continue
        name = _clean_group_name(record.get("qgroup_name") or record.get("qtype"))
        if not name:
            continue
        group_scores[name] = group_scores.get(name, 0.0) + float(record.get("full_score") or 0)
    if group_scores:
        return [{"name": name, "score": round(score, 2)} for name, score in group_scores.items()]

    rows: list[dict[str, Any]] = []
    for raw in item.get("supporting_groups") or []:
        name = _clean_group_name(raw)
        score = _score_from_group_text(raw) or _question_group_full_score(context, name)
        if name:
            rows.append({"name": name, "score": round(float(score or 0), 2)})
    return rows


def _unique_question_records(context: dict[str, Any]) -> list[dict[str, Any]]:
    seen: set[Any] = set()
    rows: list[dict[str, Any]] = []
    for record in context.get("question_records") or []:
        key = record.get("question_id") or (
            record.get("qno"),
            record.get("qgroup_name"),
            record.get("qtype"),
            record.get("co_code"),
            record.get("full_score"),
        )
        if key in seen:
            continue
        seen.add(key)
        rows.append(record)
    return rows


def _support_names_text(items: list[dict[str, Any]]) -> str:
    names = [str(item.get("name") or "").strip() for item in items if str(item.get("name") or "").strip()]
    return "、".join(names) if names else "未绑定题型"


def _clean_group_name(value: Any) -> str:
    text = re.sub(r"\s+", "", str(value or ""))
    text = re.sub(r"\d+(?:\.\d+)?分$", "", text)
    text = text.replace("单选题", "选择题")
    text = text.replace("综合应用", "综合题")
    return text


def _score_from_group_text(value: Any) -> float:
    match = re.search(r"(\d+(?:\.\d+)?)\s*分", str(value or ""))
    return float(match.group(1)) if match else 0.0


def _question_group_full_score(context: dict[str, Any], group_name: str) -> float:
    target = _clean_group_name(group_name)
    for row in context.get("question_groups") or []:
        if _clean_group_name(row.get("qgroup_name") or row.get("label")) == target:
            return float(row.get("full_score") or 0)
    total = 0.0
    for row in context.get("question_items") or []:
        if _clean_group_name(row.get("qgroup_name") or row.get("label")) == target:
            total += float(row.get("full_score") or 0)
    return total


def _support_score_count_text(context: dict[str, Any], support_items: list[dict[str, Any]]) -> str:
    parts: list[str] = []
    for item in support_items:
        name = str(item.get("name") or "")
        full_score = float(item.get("score") or 0)
        count = _count_group_scores_at_or_above(context, name, full_score * 0.6 if full_score else 0)
        if full_score:
            parts.append(f"{name}{_format_number(full_score * 0.6)}分以上的有{count}人")
    return "，".join(parts) + ("，" if parts else "")


def _count_group_scores_at_or_above(context: dict[str, Any], group_name: str, threshold: float) -> int:
    target = _clean_group_name(group_name)
    grouped: dict[str, float] = defaultdict(float)
    if context.get("question_records"):
        for record in context.get("question_records") or []:
            if _clean_group_name(record.get("qgroup_name") or record.get("qtype")) != target:
                continue
            student_no = str(record.get("student_no") or record.get("student_id") or "")
            grouped[student_no] += float(record.get("score") or 0)
    else:
        for student in context.get("students") or []:
            total = 0.0
            for label, score in (student.get("question_scores") or {}).items():
                if _clean_group_name(label).startswith(target):
                    total += float(score or 0)
            if total:
                grouped[str(student.get("student_no") or "")] = total
    return sum(1 for score in grouped.values() if score >= threshold)


def _outcome_attained_count(context: dict[str, Any], item: dict[str, Any]) -> int:
    code = str(item.get("co_code") or "")
    label = _outcome_label(item)
    threshold = float(item.get("threshold") or _default_outcome_threshold(context))
    return sum(
        1
        for row in context.get("student_outcomes") or []
        if (str(row.get("co_code") or "") == code or str(row.get("label") or "") == label)
        and float(row.get("achievement") or 0) >= threshold
    )


def _achievement_judgement(value: float, threshold: float) -> str:
    if value >= threshold + 0.05:
        return "学生整体掌握情况较好，能够较好完成该目标对应的知识与能力要求"
    if value >= threshold:
        return "学生基本达到该目标要求，但高分层次和稳定性仍有提升空间"
    return "该目标达成情况不理想，学生在相关知识理解、过程推理或综合应用方面存在明显短板"


def _outcome_label(item: dict[str, Any]) -> str:
    raw = str(item.get("label") or item.get("co_name") or item.get("co_code") or "").strip()
    number = _outcome_number(item)
    if raw.startswith("课程目标"):
        return raw
    return f"课程目标{number}" if number else raw or "课程目标"


def _outcome_number(item: dict[str, Any]) -> int:
    raw = str(item.get("label") or item.get("co_name") or item.get("co_code") or "")
    return _extract_number(r"(\d+)", raw) or 0


def _default_outcome_threshold(context: dict[str, Any]) -> float:
    outcomes = context.get("course_outcomes") or []
    if outcomes:
        return round(sum(float(item.get("threshold") or 0.65) for item in outcomes) / len(outcomes), 4)
    return 0.65


def _overall_result(context: dict[str, Any]) -> float:
    outcomes = context.get("course_outcomes") or []
    if outcomes and any(item.get("result") is not None for item in outcomes):
        return round(sum(float(item.get("result") or 0) for item in outcomes), 4)
    return float(context.get("overall_achievement") or 0)


def _add_cell_picture(cell, image_stream: BytesIO, *, width: float) -> None:
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(image_stream, width=Inches(width))


def _append_visual_appendix(document: Document, context: dict[str, Any]) -> None:
    document.add_page_break()
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading_run = heading.add_run("分析图表")
    heading_run.bold = True
    heading_run.font.size = Pt(16)

    intro = document.add_paragraph("以下图表用于辅助核对成绩分布、题型表现和课程目标达成情况。")
    intro.paragraph_format.space_after = Pt(10)

    charts = [
        ("成绩分布图", _render_score_distribution_chart(context)),
        ("课程目标达成度对比图", _render_outcome_chart(context)),
        ("题型平均得分率图", _render_question_group_chart(context)),
    ]
    if context.get("student_outcomes"):
        charts.append(("课程目标个体达成分布图", _render_student_outcome_distribution_chart(context)))
    for title, image_stream in charts:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(title)
        run.bold = True
        run.font.size = Pt(12)
        document.add_picture(image_stream, width=Inches(6.2))
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run(title).italic = True


def _configure_chart_font() -> None:
    candidates = [
        "PingFang SC",
        "Hiragino Sans GB",
        "Songti SC",
        "STHeiti",
        "Heiti TC",
        "Lantinghei SC",
        "SimHei",
        "SimSun",
        "Microsoft YaHei",
        "Arial Unicode MS",
        "Noto Sans CJK SC",
    ]
    available = {font.name for font in font_manager.fontManager.ttflist}
    for name in candidates:
        if name in available:
            plt.rcParams["font.sans-serif"] = [name]
            break
    plt.rcParams["axes.unicode_minus"] = False


def _render_score_distribution_chart(context: dict[str, Any]) -> BytesIO:
    labels = [item["label"] for item in context["score_segments"]]
    values = [item["count"] for item in context["score_segments"]]
    fig, ax = plt.subplots(figsize=(8, 4.6))
    colors = ["#1f77b4", "#4c78a8", "#72b7b2", "#54a24b", "#f58518", "#e45756"]
    ax.bar(labels, values, color=colors[: len(labels)])
    ax.set_title("成绩分布")
    ax.set_ylabel("人数")
    ax.grid(axis="y", linestyle="--", alpha=0.25)
    for index, value in enumerate(values):
        ax.text(index, value + 0.2, str(value), ha="center", va="bottom", fontsize=9)
    return _figure_to_stream(fig)


def _render_outcome_chart(context: dict[str, Any]) -> BytesIO:
    labels = [item["label"] for item in context["course_outcomes"]]
    actual = [item["achievement"] for item in context["course_outcomes"]]
    threshold = [item["threshold"] for item in context["course_outcomes"]]
    x = list(range(len(labels)))
    fig, ax = plt.subplots(figsize=(8, 4.6))
    ax.bar([item - 0.18 for item in x], actual, width=0.36, label="实际达成度", color="#2a9d8f")
    ax.bar([item + 0.18 for item in x], threshold, width=0.36, label="期望值", color="#e9c46a")
    ax.set_xticks(x)
    ax.set_xticklabels(labels)
    ax.set_ylim(0, max(actual + threshold + [0.8]) * 1.2)
    ax.set_title("课程目标达成度对比")
    ax.legend()
    ax.grid(axis="y", linestyle="--", alpha=0.25)
    return _figure_to_stream(fig)


def _render_question_group_chart(context: dict[str, Any]) -> BytesIO:
    labels = [item["qgroup_name"] for item in context["question_groups"]]
    rates = [item["achievement"] for item in context["question_groups"]]
    fig, ax = plt.subplots(figsize=(8, 4.6))
    ax.plot(labels, rates, marker="o", linewidth=2, color="#457b9d")
    ax.set_ylim(0, 1.0)
    ax.set_title("题型平均得分率")
    ax.set_ylabel("得分率")
    ax.grid(axis="y", linestyle="--", alpha=0.25)
    for index, value in enumerate(rates):
        ax.text(index, value + 0.03, _format_percent(value), ha="center", va="bottom", fontsize=9)
    return _figure_to_stream(fig)


def _render_student_outcome_distribution_chart(context: dict[str, Any]) -> BytesIO:
    rows = list(context.get("student_outcomes") or [])
    outcomes = list(context.get("course_outcomes") or [])
    targets = [
        {
            "label": str(item.get("label") or item.get("co_code") or f"课程目标{index}").strip(),
            "co_code": str(item.get("co_code") or f"CO{index}").strip(),
        }
        for index, item in enumerate(outcomes, start=1)
    ]
    if not targets:
        targets = [
            {"label": code, "co_code": code}
            for code in sorted({str(item.get("co_code") or item.get("label") or "") for item in rows if item.get("label") or item.get("co_code")})
        ]

    targets = targets[:4]
    fig, axes = plt.subplots(2, 2, figsize=(8, 5.8))
    flat_axes = axes.flatten()
    for index, ax in enumerate(flat_axes):
        if index >= len(targets):
            ax.axis("off")
            continue
        target = targets[index]
        label = target["label"]
        co_code = target["co_code"]
        matched = _student_outcome_rows_for_target(rows, label, co_code)
        values = [float(item.get("achievement") or 0) for item in matched]
        threshold = float(matched[0].get("threshold") or 0.65) if matched else 0.65
        x_values = list(range(1, len(values) + 1))
        ax.scatter(x_values, values, color="#2f80ed", s=16, label="学生达成度")
        ax.axhline(threshold, color="#d62828", linewidth=1.5, label="阈值")
        ax.set_ylim(0, 1.05)
        ax.set_title(f"{label}达成度分布", fontsize=10)
        ax.grid(axis="y", linestyle="--", alpha=0.25)
        if len(values) <= 12:
            ax.set_xticks(x_values)
        else:
            ax.set_xticks([])
        if index == 0:
            ax.legend(fontsize=8, loc="lower right")
    return _figure_to_stream(fig)


def _student_outcome_rows_for_target(rows: list[dict[str, Any]], label: str, co_code: str) -> list[dict[str, Any]]:
    label_value = str(label or "").strip()
    code_value = str(co_code or "").strip()
    return [
        item for item in rows
        if str(item.get("co_code") or "").strip() == code_value
        or str(item.get("label") or "").strip() == label_value
        or str(item.get("label") or "").strip() == code_value
    ]


def _figure_to_stream(fig) -> BytesIO:
    stream = BytesIO()
    fig.tight_layout()
    fig.savefig(stream, format="png", dpi=180, bbox_inches="tight")
    plt.close(fig)
    stream.seek(0)
    return stream


def _document_to_bytes(document: Document) -> bytes:
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output.read()


def _normalize_analysis_context(context: dict[str, Any]) -> dict[str, Any]:
    meta = dict(context.get("meta") or {})
    score_stats = dict(context.get("score_stats") or {})
    question_groups = list(context.get("question_groups") or [])
    course_outcomes = []
    for index, item in enumerate(context.get("course_outcomes") or [], start=1):
        label = str(item.get("co_name") or item.get("co_code") or f"课程目标{index}").strip() or f"课程目标{index}"
        course_outcomes.append(
            {
                "label": label,
                "co_code": item.get("co_code") or f"CO{index}",
                "co_name": label,
                "indicator": item.get("indicator") or "",
                "description": item.get("description") or "",
                "full_score": item.get("full_score") or 0,
                "avg_score": item.get("avg_score") or 0,
                "achievement": item.get("achievement") or 0,
                "threshold": item.get("threshold") or 0.65,
                "weight": item.get("weight") or 0,
                "result": item.get("result") or 0,
                "supporting_groups": item.get("supporting_groups") or [],
            }
        )

    return {
        "meta": {
            "academic_year": str(meta.get("academic_year") or ""),
            "department": str(meta.get("department") or ""),
            "course_name": str(meta.get("course_name") or ""),
            "teacher_name": str(meta.get("teacher_name") or ""),
            "class_name": str(meta.get("class_name") or ""),
            "student_count_expected": int(meta.get("student_count_expected") or score_stats.get("total_students") or 0),
            "student_count_actual": int(meta.get("student_count_actual") or score_stats.get("total_students") or 0),
            "exam_date": str(meta.get("exam_date") or ""),
        },
        "score_stats": {
            "total_students": int(score_stats.get("total_students") or len(context.get("student_scores") or [])),
            "average_score": float(score_stats.get("average_score") or 0),
            "max_score": float(score_stats.get("max_score") or 0),
            "min_score": float(score_stats.get("min_score") or 0),
            "pass_rate": float(score_stats.get("pass_rate") or 0),
        },
        "score_segments": _teacher_style_segments(context.get("student_scores") or []),
        "difficulty_label": str(context.get("difficulty_label") or "未知"),
        "question_items": list(context.get("question_items") or []),
        "question_groups": question_groups,
        "course_outcomes": course_outcomes,
        "student_outcomes": list(context.get("student_outcomes") or []),
        "question_records": list(context.get("question_records") or []),
        "overall_achievement": _analysis_overall_achievement(context.get("student_outcomes") or []),
        "warnings": list(context.get("warnings") or []),
        "narrative_sections": dict(context.get("narrative") or {}),
        "narrative": _analysis_narrative_text(context.get("narrative") or {}),
        "narrative_source": context.get("narrative_source") or "",
    }


def _teacher_style_segments(student_scores: list[dict[str, Any]]) -> list[dict[str, Any]]:
    scores = [float(item.get("final_score") or 0) for item in student_scores]
    total = len(scores) or 1
    buckets = [
        ("90—100分", lambda score: 90 <= score <= 100),
        ("80—89分", lambda score: 80 <= score < 90),
        ("70—79分", lambda score: 70 <= score < 80),
        ("60—69分", lambda score: 60 <= score < 70),
        ("50—59分", lambda score: 50 <= score < 60),
        ("50分以下", lambda score: score < 50),
    ]
    rows = []
    for label, rule in buckets:
        count = sum(1 for score in scores if rule(score))
        rows.append({"label": label, "count": count, "rate": round(count / total, 4)})
    return rows


def _analysis_overall_achievement(student_outcomes: list[dict[str, Any]]) -> float:
    grouped: dict[str, list[float]] = defaultdict(list)
    for item in student_outcomes:
        student_no = str(item.get("student_no") or "").strip()
        if not student_no:
            continue
        grouped[student_no].append(float(item.get("achievement") or 0))
    values = [sum(scores) / len(scores) for scores in grouped.values() if scores]
    if not values:
        return 0.0
    return round(sum(values) / len(values), 4)


def _analysis_narrative_text(narrative: dict[str, Any]) -> str:
    sections = {
        "score_summary": str(narrative.get("score_summary") or "").strip(),
        "support_analysis": str(narrative.get("support_analysis") or "").strip(),
        "attainment_analysis": str(narrative.get("attainment_analysis") or "").strip(),
        "improvement_actions": str(narrative.get("improvement_actions") or "").strip(),
    }
    return _combine_narrative_sections(sections)


def _combine_narrative_sections(sections: dict[str, str]) -> str:
    ordered_sections = [
        ("1、试题对课程目标的支撑度分析", str(sections.get("support_analysis") or "").strip()),
        ("2、学生作答情况对课程目标的达成度分析", str(sections.get("attainment_analysis") or "").strip()),
        ("3、教师对今后教学持续改进的具体意见", str(sections.get("improvement_actions") or "").strip()),
    ]
    lines: list[str] = []
    for title, body in ordered_sections:
        lines.append(title)
        lines.append(body or "待补充")
        lines.append("")
    return "\n".join(lines).strip()


async def enrich_context_with_ai_narrative(context: dict[str, Any]) -> dict[str, Any]:
    sections = dict(context.get("narrative_sections") or {})
    if not sections:
        sections = _fallback_sections_from_context(context)
    ai_client = DeepSeekClient()
    if ai_client.is_configured and context.get("narrative_source") != "ai_suggestion_cache":
        try:
            sections = await ai_client.generate_report_narrative(
                meta=context.get("meta") or {},
                score_stats=context.get("score_stats") or {},
                score_segments=context.get("score_segments") or [],
                question_groups=context.get("question_groups") or [],
                course_outcomes=context.get("course_outcomes") or [],
                warnings=context.get("warnings") or [],
                fallback_sections=sections,
            )
        except Exception:
            sections = sections
    enriched = dict(context)
    enriched["narrative_sections"] = sections
    enriched["narrative"] = _combine_narrative_sections(sections)
    return enriched


def _fallback_sections_from_context(context: dict[str, Any]) -> dict[str, str]:
    score_stats = context.get("score_stats") or {}
    meta = context.get("meta") or {}
    largest_segment = max(context.get("score_segments") or [{"label": "无数据", "count": 0}], key=lambda item: item["count"])
    weak_outcomes = [
        item for item in context.get("course_outcomes") or [] if float(item.get("achievement") or 0) < float(item.get("threshold") or 0.65)
    ]
    support_summary = "、".join(
        f"{item.get('co_code') or item.get('label')}达成度{_format_percent(float(item.get('achievement') or 0))}"
        for item in (context.get("course_outcomes") or [])
    ) or "当前暂无课程目标统计结果。"
    attainment_summary = (
        "当前未达标课程目标主要为"
        + "、".join(str(item.get("co_code") or item.get("label")) for item in weak_outcomes)
        + "，需要在后续教学中重点强化。"
        if weak_outcomes
        else "当前课程目标整体达到既定阈值，说明学生作答情况总体能够支撑课程目标达成。"
    )
    return {
        "score_summary": (
            f"本次《{meta.get('course_name', '')}》共纳入 {score_stats.get('total_students', 0)} 名学生，"
            f"平均分 {score_stats.get('average_score', 0)} 分，最高分 {score_stats.get('max_score', 0)} 分，"
            f"最低分 {score_stats.get('min_score', 0)} 分，及格率 {_format_percent(float(score_stats.get('pass_rate', 0)))}。"
            f"成绩主要集中在 {largest_segment.get('label', '无数据')}。"
        ),
        "support_analysis": support_summary,
        "attainment_analysis": attainment_summary,
        "improvement_actions": "建议围绕低得分题型与未达标课程目标开展针对性讲评、错题复盘和分层训练，并通过阶段性检测持续跟踪改进效果。",
    }
