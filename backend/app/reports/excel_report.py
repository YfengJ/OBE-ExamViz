from __future__ import annotations

from collections import defaultdict
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font


def build_score_report_excel(
    stats: dict,
    question_rows: list[dict],
    obe_rows: list[dict],
    warning_rows: list[dict] | None = None,
    trend_rows: list[dict] | None = None,
    meta: dict | None = None,
    sheet_name: str = "final_exam_analysis",
) -> bytes:
    output = BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        stats_df = pd.DataFrame([stats])
        question_df = pd.DataFrame(question_rows)
        obe_df = pd.DataFrame(obe_rows)
        warning_df = pd.DataFrame(warning_rows or [])
        trend_df = pd.DataFrame(trend_rows or [])
        meta_df = pd.DataFrame([meta or {}])

        meta_df = meta_df.rename(columns={
            "course_id": "课程ID", "exam_id": "考试ID", "class_name": "班级", "generated_at": "生成时间"
        })
        stats_df = stats_df.rename(columns={
            "total_students": "总人数", "average_score": "平均分", "max_score": "最高分",
            "min_score": "最低分", "pass_rate": "及格率", "score_segments": "分数段"
        })
        trend_df = trend_df.rename(columns={
            "exam_id": "考试ID", "exam_name": "考试名称", "exam_date": "考试日期", "avg_score": "平均分"
        })
        question_df = question_df.rename(columns={
            "question_id": "题目ID", "qno": "题号", "qtype": "题型", "full_score": "满分",
            "avg_score": "平均分", "difficulty": "难度", "discrimination": "区分度", "pass_rate": "得分率"
        })
        obe_df = obe_df.rename(columns={
            "co_code": "目标代码", "co_name": "课程目标", "threshold": "期望阈值",
            "full_sum": "满分总和", "actual_sum": "实际平均分", "achievement": "最终达成度"
        })
        if not warning_df.empty:
            warning_df = warning_df.rename(columns={
                "id": "记录ID", "student_id": "数据库中学生ID", "student_no": "学号", "class_name": "班级",
                "course_id": "课程ID", "course_name": "课程名称", "term": "学期", "level": "预警级别",
                "status": "处理状态", "reasons": "预警原因", "ai_summary": "AI诊断报告", "created_at": "生成时间"
            })

        meta_df.to_excel(writer, index=False, sheet_name="基础信息", startrow=0)
        stats_df.to_excel(writer, index=False, sheet_name="基础信息", startrow=3)
        trend_df.to_excel(writer, index=False, sheet_name="成绩分布")
        question_df.to_excel(writer, index=False, sheet_name="题目分析")
        obe_df.to_excel(writer, index=False, sheet_name="达成度分析")
        warning_df.to_excel(writer, index=False, sheet_name="预警记录")

    output.seek(0)
    return output.read()


def build_analysis_run_workbook(context: dict) -> bytes:
    wb = Workbook()
    ws_score = wb.active
    ws_score.title = "期末成绩单"
    ws_main = wb.create_sheet("班级分析主表")
    outcome_sheets: dict[str, any] = {}
    ws_outcomes = wb.create_sheet("课程目标平均分和达成度")
    ws_chart = wb.create_sheet("课程目标分析图")
    ws_narrative = wb.create_sheet("文字分析")

    meta = context["meta"]
    score_stats = context["score_stats"]
    student_scores = context["student_scores"]
    question_records = context["question_records"]
    question_groups = context["question_groups"]
    course_outcomes = context["course_outcomes"]
    student_outcomes = context["student_outcomes"]
    narrative = context["narrative"]

    _write_score_sheet(ws_score, meta, student_scores, question_records)
    _write_main_sheet(ws_main, meta, score_stats, student_scores, student_outcomes)
    for item in course_outcomes:
        sheet_title = f"{meta['class_name']}班{str(item['co_code']).replace('CO', '目标')}"[:31]
        ws = wb.create_sheet(sheet_title)
        outcome_sheets[item["co_code"]] = ws
        _write_single_outcome_sheet(ws, item, student_outcomes)
    _write_outcome_summary_sheet(ws_outcomes, course_outcomes)
    _write_chart_sheet(ws_chart, course_outcomes)
    _write_narrative_sheet(ws_narrative, meta, score_stats, context["difficulty_label"], narrative)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


def _write_score_sheet(ws, meta: dict, student_scores: list[dict], question_records: list[dict]) -> None:
    question_df = pd.DataFrame(question_records)
    group_order = []
    if not question_df.empty:
        group_order = question_df[["qgroup_name", "qno"]].drop_duplicates().to_dict(orient="records")
    headers = ["学号", "姓名"] + [f"{item['qgroup_name']}-{item['qno']}" for item in group_order] + ["期末总分"]
    ws.append([f"{meta['academic_year']}学年 {meta['term_label']} 期末成绩单"])
    ws.append([f"课程：{meta['course_name']}", "", f"班级：{meta['class_name']}", "", f"教师：{meta['teacher_name']}"])
    ws.append([f"考试日期：{meta['exam_date']}", "", f"应考/实考：{meta['student_count_expected']}/{meta['student_count_actual']}"])
    ws.append([])
    ws.append(headers)
    question_map = defaultdict(dict)
    if not question_df.empty:
        for row in question_df.to_dict(orient="records"):
            question_map[row["student_no"]][f"{row['qgroup_name']}-{row['qno']}"] = row["score"]
    for row in student_scores:
        ws.append(
            [row["student_no"], row.get("name", "")]
            + [question_map[row["student_no"]].get(header, 0) for header in headers[2:-1]]
            + [row.get("final_score", 0)]
        )
    _beautify_sheet(ws)


def _write_main_sheet(ws, meta: dict, score_stats: dict, student_scores: list[dict], student_outcomes: list[dict]) -> None:
    outcome_codes = sorted({item["co_code"] for item in student_outcomes})
    headers = ["学号", "姓名", "平时成绩", "期中成绩", "期末成绩", "课程总评"]
    for code in outcome_codes:
        headers.extend([f"{code}成绩", f"{code}达成度"])
    headers.append("课程目标总达成度")
    ws.append(["试卷分析主表"])
    ws.append([f"课程：{meta['course_name']}", "", f"班级：{meta['class_name']}", "", f"平均分：{score_stats['average_score']}", "", f"及格率：{round(score_stats['pass_rate'] * 100, 2)}%"])
    ws.append([])
    ws.append(headers)
    grouped = defaultdict(dict)
    for row in student_outcomes:
        grouped[row["student_no"]][row["co_code"]] = row
    for row in student_scores:
        line = [
            row["student_no"],
            row.get("name", ""),
            row.get("usual_score", 0),
            row.get("midterm_score", 0),
            row.get("final_score", 0),
            row.get("course_total_score", 0),
        ]
        achievements = []
        for code in outcome_codes:
            item = grouped[row["student_no"]].get(code, {})
            line.append(item.get("co_score", 0))
            line.append(item.get("achievement", 0))
            achievements.append(item.get("achievement", 0))
        line.append(round(sum(achievements) / len(achievements), 4) if achievements else 0)
        ws.append(line)
    _beautify_sheet(ws)


def _write_single_outcome_sheet(ws, outcome: dict, student_outcomes: list[dict]) -> None:
    ws.append([f"{outcome['co_code']} 达成明细"])
    ws.append(["学号", "姓名", "课程目标得分", "满分", "达成度", "阈值", "是否达成"])
    for row in [item for item in student_outcomes if item["co_code"] == outcome["co_code"]]:
        ws.append(
            [
                row["student_no"],
                row["student_name"],
                row["co_score"],
                row["full_score"],
                row["achievement"],
                row["threshold"],
                "达成" if row["is_attained"] else "未达成",
            ]
        )
    _beautify_sheet(ws)


def _write_outcome_summary_sheet(ws, course_outcomes: list[dict]) -> None:
    ws.append(["课程目标", "总分", "平均分", "分项达成度", "阈值", "权重", "达成结果"])
    for row in course_outcomes:
        ws.append(
            [
                row["co_code"],
                row["full_score"],
                row["avg_score"],
                row["achievement"],
                row["threshold"],
                row["weight"],
                row["result"],
            ]
        )
    _beautify_sheet(ws)


def _write_chart_sheet(ws, course_outcomes: list[dict]) -> None:
    ws.append(["课程目标", "预测值", "真实值"])
    for row in course_outcomes:
        ws.append([row["co_code"], row["threshold"], row["achievement"]])
    chart = BarChart()
    chart.title = "课程目标达成度对比"
    chart.y_axis.title = "达成度"
    chart.x_axis.title = "课程目标"
    data = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=len(course_outcomes) + 1)
    categories = Reference(ws, min_col=1, min_row=2, max_row=len(course_outcomes) + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 8
    chart.width = 14
    ws.add_chart(chart, "E2")
    _beautify_sheet(ws)


def _write_narrative_sheet(ws, meta: dict, score_stats: dict, difficulty_label: str, narrative: dict) -> None:
    ws.append([f"{meta['department']} {meta['academic_year']} {meta['term_label']} 试卷分析文字稿"])
    ws.append([f"课程：{meta['course_name']}"])
    ws.append([f"班级：{meta['class_name']}"])
    ws.append([f"平均分：{score_stats['average_score']}，最高分：{score_stats['max_score']}，最低分：{score_stats['min_score']}，难度：{difficulty_label}"])
    ws.append([])
    ws.append(["成绩统计摘要", narrative["score_summary"]])
    ws.append(["课程目标支撑度分析", narrative["support_analysis"]])
    ws.append(["课程目标达成度分析", narrative["attainment_analysis"]])
    ws.append(["持续改进建议", narrative["improvement_actions"]])
    _beautify_sheet(ws)


def _beautify_sheet(ws) -> None:
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    for cell in ws[1]:
        cell.font = Font(bold=True, size=13)
    for column_cells in ws.columns:
        length = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 12), 28)
